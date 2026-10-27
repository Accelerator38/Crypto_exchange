from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import TwoSlopeNorm
import numpy as np
from openpyxl import load_workbook
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from exia.genetic_primus.contracts import STATIC_FEATURE_NAMES, load_primus_config, sha256_file  # noqa: E402
from exia.genetic_primus.data import load_panel  # noqa: E402
from exia.genetic_primus.features import RobustScaler, build_feature_frame  # noqa: E402
from exia.genetic_primus.policy import GeneticPrimusPolicy  # noqa: E402
from exia.genetic_primus.simulator import simulate_policy  # noqa: E402
from exia.scoring import moving_block_lcb  # noqa: E402


DEFAULT_CONFIG = ROOT / "configs/genetic_primus_v3.json"
DEFAULT_RUN = ROOT / "Reports/Exia/Genetic_Primus/development_v3_100g_20260821"
DEFAULT_OUTPUT = ROOT / "Reports/Exia/Genetic_Primus/v3_generalization_audit_20260824"
ROUND_TRIP_COST_BPS = 16.0


def population_stability_index(reference: pd.Series, observed: pd.Series, bins: int = 10) -> float:
    ref = pd.to_numeric(reference, errors="coerce").dropna().to_numpy(dtype=float)
    obs = pd.to_numeric(observed, errors="coerce").dropna().to_numpy(dtype=float)
    if len(ref) < bins * 2 or len(obs) < bins * 2:
        return float("nan")
    edges = np.unique(np.quantile(ref, np.linspace(0.0, 1.0, bins + 1)))
    if len(edges) < 3:
        return 0.0
    edges[0] = -np.inf
    edges[-1] = np.inf
    ref_counts = np.histogram(ref, bins=edges)[0].astype(float)
    obs_counts = np.histogram(obs, bins=edges)[0].astype(float)
    epsilon = 1.0e-6
    ref_share = np.maximum(ref_counts / ref_counts.sum(), epsilon)
    obs_share = np.maximum(obs_counts / obs_counts.sum(), epsilon)
    return float(np.sum((obs_share - ref_share) * np.log(obs_share / ref_share)))


def _load_scaler(path: Path) -> RobustScaler:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return RobustScaler(
        median=np.asarray(payload["median"], dtype=np.float64),
        scale=np.asarray(payload["scale"], dtype=np.float64),
    )


def build_drift_table(features: pd.DataFrame) -> pd.DataFrame:
    local = features.copy()
    local["year"] = pd.to_datetime(local["timestamp"], unit="ms", utc=True).dt.year
    reference = local.loc[local["year"].isin([2022, 2023])]
    observed = local.loc[local["year"] == 2024]
    rows: list[dict[str, Any]] = []
    for feature in STATIC_FEATURE_NAMES:
        ref = reference[feature].astype(float)
        obs = observed[feature].astype(float)
        ref_iqr = float(ref.quantile(0.75) - ref.quantile(0.25))
        median_shift = float(obs.median() - ref.median())
        standardized = median_shift / ref_iqr if abs(ref_iqr) > 1.0e-12 else 0.0
        psi = population_stability_index(ref, obs)
        rows.append(
            {
                "feature": feature,
                "reference_period": "2022-2023",
                "observed_period": "2024",
                "reference_median": float(ref.median()),
                "observed_median": float(obs.median()),
                "reference_iqr": ref_iqr,
                "standardized_median_shift": standardized,
                "psi": psi,
                "reference_missing_rate": float(ref.isna().mean()),
                "observed_missing_rate": float(obs.isna().mean()),
                "drift_level": "HIGH" if psi >= 0.25 else "MODERATE" if psi >= 0.10 else "LOW",
            }
        )
    return pd.DataFrame(rows).sort_values(["psi", "feature"], ascending=[False, True])


def build_counterfactual_ledger(
    features: pd.DataFrame,
    entries: pd.DataFrame,
    *,
    cost_bps: float,
) -> pd.DataFrame:
    groups = {
        str(symbol): group.sort_values("timestamp", kind="stable").reset_index(drop=True)
        for symbol, group in features.groupby("symbol", sort=True, observed=True)
    }
    index_maps = {
        symbol: {int(timestamp): index for index, timestamp in enumerate(group["timestamp"])}
        for symbol, group in groups.items()
    }
    rows: list[dict[str, Any]] = []
    for trade in entries.itertuples(index=False):
        symbol = str(trade.symbol)
        group = groups[symbol]
        entry_index = index_maps[symbol].get(int(trade.entry_timestamp))
        if entry_index is None:
            continue
        direction = 1 if str(trade.direction) == "LONG" else -1
        entry_price = float(trade.entry_price)
        candidates: list[tuple[str, int, str]] = [
            ("fixed_h1", entry_index + 1, "fixed_horizon"),
            ("fixed_h3", entry_index + 3, "fixed_horizon"),
            ("fixed_h6", entry_index + 6, "fixed_horizon"),
        ]
        degradation_index = min(entry_index + 12, len(group) - 1)
        for index in range(entry_index, min(entry_index + 12, len(group) - 1)):
            if str(group.iloc[index]["market_state"]) != str(trade.market_state):
                degradation_index = index + 1
                break
        candidates.append(("regime_degradation_h12", degradation_index, "regime_change_or_cap"))
        for policy_name, exit_index, reason in candidates:
            if exit_index >= len(group):
                continue
            exit_row = group.iloc[exit_index]
            exit_price = float(exit_row["open"])
            gross = direction * (exit_price / entry_price - 1.0) * 10_000.0
            rows.append(
                {
                    "symbol": symbol,
                    "direction": str(trade.direction),
                    "market_state": str(trade.market_state),
                    "signal_timestamp": int(trade.signal_timestamp),
                    "entry_timestamp": int(trade.entry_timestamp),
                    "exit_timestamp": int(exit_row["timestamp"]),
                    "exit_policy": policy_name,
                    "exit_reason": reason,
                    "holding_bars": int(exit_index - entry_index),
                    "gross_bps": gross,
                    "cost_bps": float(cost_bps),
                    "net_bps": gross - float(cost_bps),
                    "fills": 2,
                }
            )
    return pd.DataFrame(rows)


def summarize_exits(ledger: pd.DataFrame) -> pd.DataFrame:
    local = ledger.copy()
    local["year"] = pd.to_datetime(local["entry_timestamp"], unit="ms", utc=True).dt.year
    rows: list[dict[str, Any]] = []
    for (year, policy), subset in local.groupby(["year", "exit_policy"], sort=True):
        values = subset["net_bps"].astype(float)
        rows.append(
            {
                "year": int(year),
                "exit_policy": str(policy),
                "trades": int(len(subset)),
                "mean_bps": float(values.mean()),
                "median_bps": float(values.median()),
                "lcb_bps": moving_block_lcb(
                    values,
                    subset["signal_timestamp"],
                    samples=2_000,
                    alpha=0.05,
                    seed=int.from_bytes(
                        hashlib.sha256(f"v3-audit:{year}:{policy}".encode()).digest()[:8], "big"
                    ),
                ),
                "positive_rate": float((values > 0.0).mean()),
                "mean_holding_bars": float(subset["holding_bars"].mean()),
            }
        )
    return pd.DataFrame(rows)


def _write_sheet(writer: pd.ExcelWriter, name: str, frame: pd.DataFrame) -> None:
    frame.to_excel(writer, sheet_name=name, index=False)
    sheet = writer.book[name]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        values = [str(cell.value) if cell.value is not None else "" for cell in column[:250]]
        sheet.column_dimensions[column[0].column_letter].width = min(
            48, max(10, max(len(value) for value in values) + 2)
        )


def write_workbook(path: Path, frames: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in frames.items():
            _write_sheet(writer, name, frame)
    workbook = load_workbook(path, read_only=False)
    for name, frame in frames.items():
        sheet = workbook[name]
        if sheet.max_column != len(frame.columns) or sheet.max_column <= 1:
            raise RuntimeError(f"invalid separate-cell layout: {name}")
        if sheet.merged_cells.ranges:
            raise RuntimeError(f"merged cells are forbidden: {name}")
    workbook.close()


def _plot_drift(drift: pd.DataFrame, path: Path) -> None:
    ordered = drift.sort_values("psi")
    colors = ["#B42318" if value >= 0.25 else "#D97706" if value >= 0.1 else "#2878B5" for value in ordered["psi"]]
    figure, axis = plt.subplots(figsize=(12, 9), constrained_layout=True)
    axis.barh(ordered["feature"], ordered["psi"], color=colors)
    axis.axvline(0.10, color="#D97706", linestyle="--", linewidth=1)
    axis.axvline(0.25, color="#B42318", linestyle="--", linewidth=1)
    axis.set_xlabel("Population Stability Index, 2024 vs 2022-2023")
    axis.set_title("Genetic_Primus V3 feature drift")
    axis.grid(axis="x", alpha=0.2)
    figure.savefig(path, dpi=180)
    plt.close(figure)


def _plot_exits(summary: pd.DataFrame, path: Path) -> None:
    matrix = summary.pivot(index="exit_policy", columns="year", values="lcb_bps").reindex(columns=[2022, 2023, 2024])
    values = matrix.to_numpy(dtype=float)
    finite = np.abs(values[np.isfinite(values)])
    limit = max(1.0, float(np.quantile(finite, 0.95))) if len(finite) else 1.0
    figure, axis = plt.subplots(figsize=(9, 5), constrained_layout=True)
    image = axis.imshow(values, cmap="RdYlGn", norm=TwoSlopeNorm(vmin=-limit, vcenter=0.0, vmax=limit), aspect="auto")
    axis.set_xticks(range(len(matrix.columns)), matrix.columns)
    axis.set_yticks(range(len(matrix.index)), matrix.index)
    axis.set_title("Entry feasibility by exit policy: costed LCB")
    for row in range(values.shape[0]):
        for column in range(values.shape[1]):
            value = values[row, column]
            axis.text(column, row, "n/a" if not np.isfinite(value) else f"{value:.1f}", ha="center", va="center")
    figure.colorbar(image, ax=axis, label="bps/trade")
    figure.savefig(path, dpi=180)
    plt.close(figure)


def run(config_path: Path, run_dir: Path, output_dir: Path) -> dict[str, Any]:
    terminal = json.loads((run_dir / "terminal_rejection_manifest.json").read_text(encoding="utf-8"))
    if terminal.get("status") != "TERMINAL_REJECTED_VALIDATION_GENERALIZATION":
        raise ValueError("V3 audit requires terminal rejection")
    config = load_primus_config(config_path, root=ROOT)
    panel = load_panel(config, root=ROOT)
    features = build_feature_frame(panel)
    drift = build_drift_table(features)
    scaler = _load_scaler(run_dir / "scaler.json")
    with np.load(run_dir / "champion.npz", allow_pickle=False) as payload:
        genome = np.asarray(payload["genome"], dtype=np.float32)
    simulation = simulate_policy(
        features,
        policy=GeneticPrimusPolicy(genome),
        scaler=scaler,
        start_timestamp=int(pd.Timestamp("2022-01-01T00:00:00Z").timestamp() * 1000),
        end_timestamp=int(pd.Timestamp("2025-01-01T00:00:00Z").timestamp() * 1000),
        round_trip_cost_bps=ROUND_TRIP_COST_BPS,
    )
    neural = simulation.ledger.copy()
    neural["exit_policy"] = "neural_dynamic"
    counterfactual = build_counterfactual_ledger(
        features,
        neural,
        cost_bps=ROUND_TRIP_COST_BPS,
    )
    exit_ledger = pd.concat([neural, counterfactual], ignore_index=True, sort=False)
    exit_summary = summarize_exits(exit_ledger)
    output_dir.mkdir(parents=True, exist_ok=True)
    drift.to_parquet(output_dir / "feature_drift.parquet", index=False)
    exit_ledger.to_parquet(output_dir / "entry_exit_ledger.parquet", index=False)
    exit_summary.to_parquet(output_dir / "entry_exit_summary.parquet", index=False)
    write_workbook(
        output_dir / "v3_generalization_audit.xlsx",
        {"Feature Drift": drift, "Exit Summary": exit_summary, "Exit Ledger": exit_ledger},
    )
    _plot_drift(drift, output_dir / "feature_drift.png")
    _plot_exits(exit_summary, output_dir / "entry_exit_lcb.png")
    high_drift = drift.loc[drift["drift_level"] == "HIGH", "feature"].tolist()
    year_2024 = exit_summary.loc[exit_summary["year"] == 2024].sort_values("lcb_bps", ascending=False)
    best_exit = year_2024.iloc[0]
    report = [
        "# Genetic_Primus V3 generalization audit",
        "",
        "Status: diagnostic only; V3 remains terminally rejected.",
        "",
        "## Feature drift",
        "",
        f"- High-PSI features: {', '.join(high_drift) or 'none'}.",
        f"- High/moderate/low counts: {(drift.drift_level == 'HIGH').sum()}/{(drift.drift_level == 'MODERATE').sum()}/{(drift.drift_level == 'LOW').sum()}.",
        "",
        "## Entry/exit counterfactual",
        "",
        "| Year | Exit | Trades | Mean | Median | LCB | Hold bars |",
        "|---:|---|---:|---:|---:|---:|---:|",
    ]
    for row in exit_summary.itertuples(index=False):
        report.append(
            f"| {row.year} | {row.exit_policy} | {row.trades} | {row.mean_bps:.2f} | {row.median_bps:.2f} | {row.lcb_bps:.2f} | {row.mean_holding_bars:.2f} |"
        )
    report.extend(
        [
            "",
            f"Best 2024 counterfactual by LCB: `{best_exit.exit_policy}`, LCB {best_exit.lcb_bps:.2f} bps.",
            "A nonpositive best LCB means that changing only the exit cannot repair V3 entries.",
            "",
            "No OOS, paper or live stage was opened.",
            "",
        ]
    )
    (output_dir / "report.md").write_text("\n".join(report), encoding="utf-8")
    manifest = {
        "schema_version": "exia.genetic_primus_generalization_audit.v1",
        "generated_at": datetime.now(UTC).isoformat(),
        "v3_terminal_manifest_sha256": sha256_file(run_dir / "terminal_rejection_manifest.json"),
        "champion_genome_sha256": hashlib.sha256(genome.tobytes()).hexdigest(),
        "dataset_sha256": config.dataset["dataset_sha256"],
        "years": [2022, 2023, 2024],
        "round_trip_cost_bps": ROUND_TRIP_COST_BPS,
        "high_drift_features": high_drift,
        "best_2024_exit": str(best_exit.exit_policy),
        "best_2024_exit_lcb_bps": float(best_exit.lcb_bps),
        "oos_opened": False,
        "paper_allowed": False,
        "live_allowed": False,
        "orders_enabled": False,
        "promotion_authority": False,
    }
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description="Audit V3 drift and entry/exit generalization")
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--run-dir", type=Path, default=DEFAULT_RUN)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = run(args.config.resolve(), args.run_dir.resolve(), args.output.resolve())
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
