from __future__ import annotations

import argparse
import json
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from exia.genetic_primus.contracts import sha256_file  # noqa: E402
from exia.genetic_primus.contracts_v5 import load_v5_contract  # noqa: E402
from exia.genetic_primus.data_source_v5 import validate_dataset_manifest  # noqa: E402


DEFAULT_CONTRACT = ROOT / "configs/genetic_primus_v5_contract.json"
DEFAULT_DATASET_MANIFEST = (
    ROOT / "Retrodate/genetic_primus_v5_binance_metrics_v1_20220101_20241231/manifest.json"
)
DEFAULT_OUTPUT = ROOT / "Reports/Exia/Genetic_Primus/v5_data_source_gate_20260824"


def _workbook(path: Path, frames: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.book[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
    workbook = load_workbook(path, read_only=False)
    for name, frame in frames.items():
        sheet = workbook[name]
        if sheet.max_column != len(frame.columns) or sheet.max_column <= 1 or sheet.merged_cells.ranges:
            raise RuntimeError(f"invalid V5 data gate workbook: {name}")
    workbook.close()


def evaluate(contract_path: Path, dataset_manifest_path: Path, output_dir: Path) -> dict[str, Any]:
    contract = load_v5_contract(contract_path, root=ROOT)
    dataset_manifest = validate_dataset_manifest(contract, dataset_manifest_path, root=ROOT)
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / "report.json"
    if report_path.is_file():
        existing = json.loads(report_path.read_text(encoding="utf-8"))
        if existing.get("contract_sha256") == contract.sha256 and existing.get(
            "dataset_manifest_sha256"
        ) == sha256_file(dataset_manifest_path):
            return existing
        raise ValueError("stale V5 data-source gate report")
    checks: list[dict[str, Any]] = []

    def add(check_id: str, actual: Any, required: Any, passed: bool) -> None:
        checks.append({"check_id": check_id, "actual": actual, "required": required, "passed": bool(passed)})

    source = contract.source
    add(
        "new_source_family",
        dataset_manifest["source_family"],
        "not present in prior_terminal_experiments",
        dataset_manifest["source_family"] not in set(contract.payload["prior_terminal_experiments"]),
    )
    add("point_in_time", dataset_manifest["point_in_time"], True, dataset_manifest["point_in_time"] is True)
    add(
        "official_checksums_verified",
        dataset_manifest["official_checksums_verified"],
        True,
        dataset_manifest["official_checksums_verified"] is True,
    )
    add(
        "raw_archive_retained",
        dataset_manifest["raw_archive_retained"],
        True,
        dataset_manifest["raw_archive_retained"] is True,
    )
    add(
        "raw_file_coverage",
        float(dataset_manifest["raw_file_coverage"]),
        float(source["minimum_raw_file_coverage"]),
        float(dataset_manifest["raw_file_coverage"]) >= float(source["minimum_raw_file_coverage"]),
    )
    add(
        "symbol_count",
        len(dataset_manifest["coverage"]),
        int(source["minimum_symbols"]),
        len(dataset_manifest["coverage"]) >= int(source["minimum_symbols"]),
    )
    coverage_rows: list[dict[str, Any]] = []
    for row in dataset_manifest["coverage"]:
        symbol = str(row["symbol"])
        add(
            f"{symbol}:bar_coverage",
            float(row["bar_coverage"]),
            float(source["minimum_complete_bar_coverage"]),
            float(row["bar_coverage"]) >= float(source["minimum_complete_bar_coverage"]),
        )
        add(
            f"{symbol}:complete_bar_rate",
            float(row["complete_bar_rate"]),
            float(source["minimum_complete_bar_coverage"]),
            float(row["complete_bar_rate"]) >= float(source["minimum_complete_bar_coverage"]),
        )
        for feature in source["required_core_features"]:
            key = f"{feature}_coverage"
            value = float(row[key])
            passed = value >= float(source["minimum_core_feature_coverage"])
            add(
                f"{symbol}:{feature}",
                value,
                float(source["minimum_core_feature_coverage"]),
                passed,
            )
            coverage_rows.append(
                {
                    "symbol": symbol,
                    "feature": feature,
                    "coverage": value,
                    "required": float(source["minimum_core_feature_coverage"]),
                    "passed": passed,
                }
            )
    failed = [item["check_id"] for item in checks if not item["passed"]]
    accepted = not failed
    if accepted:
        status = "V5_DATA_INTEGRITY_ACCEPTED_FEASIBILITY_PENDING"
        terminal_reason = None
    else:
        status = "V5_TERMINAL_REJECTED_DATA_SOURCE_INSUFFICIENT_COVERAGE"
        terminal_reason = "DATA_SOURCE_INSUFFICIENT_COVERAGE"
    report = {
        "schema_version": "exia.genetic_primus_data_source_gate.v5",
        "generated_at": datetime.now(UTC).isoformat(),
        "status": status,
        "terminal_reason": terminal_reason,
        "contract_sha256": contract.sha256,
        "dataset_manifest_sha256": sha256_file(dataset_manifest_path),
        "dataset_sha256": dataset_manifest["dataset_sha256"],
        "source_id": dataset_manifest["source_id"],
        "checks": checks,
        "failed_checks": failed,
        "integrity_and_coverage_passed": accepted,
        "feasibility_labels_opened": False,
        "evolution_allowed": False,
        "validation_opened": False,
        "oos_opened": False,
        "paper_allowed": False,
        "live_allowed": False,
        "orders_enabled": False,
        "promotion_authority": False,
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    checks_frame = pd.DataFrame(checks)
    coverage_frame = pd.DataFrame(coverage_rows)
    _workbook(output_dir / "data_source_gate.xlsx", {"Checks": checks_frame, "Feature Coverage": coverage_frame})
    failed_taker = coverage_frame.loc[
        (coverage_frame["feature"] == "taker_imbalance_mean") & (~coverage_frame["passed"])
    ]
    markdown = [
        "# Genetic_Primus V5 data-source acceptance",
        "",
        f"Status: `{status}`.",
        "",
        f"Raw files: {dataset_manifest['downloaded_raw_files']}/{dataset_manifest['expected_raw_files']}.",
        f"Raw rows: {dataset_manifest['raw_rows']}; derived 4h rows: {dataset_manifest['derived_rows']}.",
        f"Official files with repaired row ordering: {dataset_manifest['timestamp_order_repaired_files']}.",
        "",
    ]
    if not failed_taker.empty:
        markdown.extend(
            [
                "The preregistered 90% core-feature coverage gate failed for taker imbalance:",
                "",
                *[
                    f"- {row.symbol}: {row.coverage:.4%}."
                    for row in failed_taker.itertuples(index=False)
                ],
                "",
                "The experiment contract is immutable. Thresholds and required features cannot be changed after observing coverage.",
                "Feasibility labels remain unopened and neuroevolution is terminally blocked for V5.",
            ]
        )
    (output_dir / "report.md").write_text("\n".join(markdown) + "\n", encoding="utf-8")
    freeze = {
        "schema_version": "exia.genetic_primus_terminal_freeze.v5",
        "generated_at": report["generated_at"],
        "status": status,
        "terminal_reason": terminal_reason,
        "contract_sha256": contract.sha256,
        "dataset_manifest_sha256": report["dataset_manifest_sha256"],
        "dataset_sha256": report["dataset_sha256"],
        "gate_report_sha256": sha256_file(report_path),
        "evolution_allowed": False,
        "validation_allowed": False,
        "oos_allowed": False,
        "paper_allowed": False,
        "live_allowed": False,
        "orders_enabled": False,
        "promotion_authority": False,
    }
    (output_dir / "terminal_freeze.json").write_text(
        json.dumps(freeze, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return report


def assert_evolution_authority(gate_report_path: Path) -> dict[str, Any]:
    report = json.loads(gate_report_path.read_text(encoding="utf-8"))
    if report.get("status") != "V5_DATA_SOURCE_ACCEPTED" or report.get("evolution_allowed") is not True:
        raise PermissionError(f"Genetic_Primus V5 evolution blocked by data-source gate: {report.get('status')}")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Evaluate the immutable Genetic_Primus V5 data-source gate")
    parser.add_argument("--contract", type=Path, default=DEFAULT_CONTRACT)
    parser.add_argument("--dataset-manifest", type=Path, default=DEFAULT_DATASET_MANIFEST)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = evaluate(args.contract.resolve(), args.dataset_manifest.resolve(), args.output_dir.resolve())
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
