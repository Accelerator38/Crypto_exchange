from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
import time
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import numpy as np
from openpyxl import load_workbook
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from exia.genetic_primus.contracts import sha256_file  # noqa: E402
from exia.genetic_primus.contracts_v7 import load_v7_contract  # noqa: E402
from exia.genetic_primus.data_source_v5 import pack_raw_directory, parse_checksum, sha256_bytes  # noqa: E402
from exia.genetic_primus.data_source_v7 import (  # noqa: E402
    aggregate_liquidations,
    four_hour_grid,
    read_liquidation_zip,
)


DEFAULT_CONTRACT = ROOT / "configs/genetic_primus_v7_contract.json"
DEFAULT_OUTPUT = ROOT / "Retrodate/genetic_primus_v7_coinm_liquidations_v1_20230625_20241014"
S3_NAMESPACE = {"s3": "http://s3.amazonaws.com/doc/2006-03-01/"}


def _fetch(url: str, attempts: int = 6) -> bytes:
    error: Exception | None = None
    for attempt in range(attempts):
        try:
            request = Request(url, headers={"User-Agent": "Exia-V7-Research/1.0"})
            with urlopen(request, timeout=60) as response:
                return response.read()
        except (HTTPError, URLError, TimeoutError) as exc:
            error = exc
            if isinstance(exc, HTTPError) and exc.code == 404:
                break
            time.sleep(min(10.0, 0.5 * (2**attempt)))
    raise RuntimeError(f"failed to fetch {url}: {error}")


def _list_source(contract) -> pd.DataFrame:
    source = contract.liquidation_source
    start = pd.Timestamp(source["start_date"], tz="UTC")
    end = pd.Timestamp(source["end_date"], tz="UTC")
    rows: list[dict[str, Any]] = []
    for target_symbol, item in source["symbols"].items():
        exchange_symbol = item["exchange_symbol"]
        url = source["s3_listing_url"].format(exchange_symbol=exchange_symbol)
        root = ET.fromstring(_fetch(url))
        if root.findtext("s3:IsTruncated", default="false", namespaces=S3_NAMESPACE).lower() != "false":
            raise ValueError(f"V7 listing unexpectedly truncated: {exchange_symbol}")
        for content in root.findall("s3:Contents", S3_NAMESPACE):
            key = content.findtext("s3:Key", default="", namespaces=S3_NAMESPACE)
            if not key.endswith(".zip"):
                continue
            date_text = key.removesuffix(".zip").rsplit("-", 3)[-3:]
            date = "-".join(date_text)
            try:
                timestamp = pd.Timestamp(date, tz="UTC")
            except ValueError:
                continue
            if not start <= timestamp <= end:
                continue
            filename = key.rsplit("/", 1)[-1]
            rows.append(
                {
                    "symbol": target_symbol,
                    "exchange_symbol": exchange_symbol,
                    "date": date,
                    "filename": filename,
                    "key": key,
                    "bytes_listed": int(content.findtext("s3:Size", default="0", namespaces=S3_NAMESPACE)),
                    "etag": content.findtext("s3:ETag", default="", namespaces=S3_NAMESPACE).strip('"'),
                    "last_modified": content.findtext("s3:LastModified", default="", namespaces=S3_NAMESPACE),
                    "url": source["download_url"].format(exchange_symbol=exchange_symbol, date=date),
                }
            )
    listing = pd.DataFrame(rows).sort_values(["exchange_symbol", "date"], kind="stable").reset_index(drop=True)
    if listing.empty or listing.duplicated(["exchange_symbol", "date"]).any():
        raise ValueError("V7 source listing is empty or duplicated")
    return listing


def _download(item: dict[str, Any], raw_dir: Path) -> dict[str, Any]:
    target = raw_dir / str(item["exchange_symbol"]) / str(item["filename"])
    target.parent.mkdir(parents=True, exist_ok=True)
    official = parse_checksum(_fetch(str(item["url"]) + ".CHECKSUM"), str(item["filename"]))
    if not target.is_file() or sha256_file(target) != official:
        payload = _fetch(str(item["url"]))
        if sha256_bytes(payload) != official:
            raise ValueError(f"V7 official checksum mismatch: {item['filename']}")
        temporary = target.with_suffix(".zip.tmp")
        temporary.write_bytes(payload)
        temporary.replace(target)
    return {
        **item,
        "local_path": target.relative_to(raw_dir).as_posix(),
        "bytes": target.stat().st_size,
        "sha256": official,
        "official_sha256": official,
        "downloaded": True,
        "error": None,
    }


def _workbook(path: Path, frames: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.book[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
    workbook = load_workbook(path, read_only=False)
    for name, frame in frames.items():
        sheet = workbook[name]
        if sheet.max_column != len(frame.columns) or sheet.max_column <= 1 or sheet.merged_cells.ranges:
            raise RuntimeError(f"invalid V7 dataset workbook: {name}")
    workbook.close()


def _build_panels(contract, catalog: pd.DataFrame, raw_dir: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    source = contract.liquidation_source
    days = pd.date_range(source["start_date"], source["end_date"], freq="1D", tz="UTC")
    interval_ms = 14_400_000
    panel_parts: list[pd.DataFrame] = []
    quality_rows: list[dict[str, Any]] = []
    for target_symbol, mapping in source["symbols"].items():
        subset = catalog.loc[(catalog["symbol"] == target_symbol) & (catalog["downloaded"] == True)].copy()  # noqa: E712
        published_dates = set(subset["date"].astype(str))
        aggregated: list[pd.DataFrame] = []
        raw_rows = 0
        deduplicated_rows = 0
        for row in subset.itertuples(index=False):
            frame = read_liquidation_zip(raw_dir / str(row.local_path), expected_date=str(row.date))
            raw_rows += len(frame) + int(frame.attrs["exact_duplicates_removed"])
            deduplicated_rows += int(frame.attrs["exact_duplicates_removed"])
            if not frame.empty:
                aggregated.append(
                    aggregate_liquidations(
                        frame,
                        target_symbol=target_symbol,
                        source_symbol=str(mapping["exchange_symbol"]),
                        contract_size_usd=float(mapping["contract_size_usd"]),
                    )
                )
        grid = four_hour_grid(source["start_date"], source["end_date"])
        grid["archive_available"] = grid["date"].isin(published_dates)
        grid["symbol"] = target_symbol
        grid["source_symbol"] = str(mapping["exchange_symbol"])
        bars = pd.concat(aggregated, ignore_index=True) if aggregated else pd.DataFrame()
        if not bars.empty and bars.duplicated(["timestamp", "symbol"]).any():
            raise ValueError(f"duplicate V7 aggregated bars: {target_symbol}")
        value_columns = [
            "sell_liquidation_usd",
            "buy_liquidation_usd",
            "sell_liquidation_count",
            "buy_liquidation_count",
            "liquidation_total_usd",
        ]
        if bars.empty:
            for column in value_columns:
                grid[column] = np.nan
        else:
            grid = grid.merge(
                bars.drop(columns=["symbol", "source_symbol"]),
                on="timestamp",
                how="left",
                validate="one_to_one",
            )
        for column in value_columns:
            grid.loc[grid["archive_available"] & grid[column].isna(), column] = 0.0
            grid.loc[~grid["archive_available"], column] = np.nan
        panel_parts.append(grid)
        quality_rows.append(
            {
                "symbol": target_symbol,
                "exchange_symbol": mapping["exchange_symbol"],
                "expected_days": len(days),
                "published_days": len(published_dates),
                "archive_day_coverage": len(published_dates) / len(days),
                "raw_rows": raw_rows,
                "exact_duplicates_removed": deduplicated_rows,
                "deduplicated_rows": raw_rows - deduplicated_rows,
            }
        )
    panel = pd.concat(panel_parts, ignore_index=True).sort_values(["timestamp", "symbol"], kind="stable")
    anchor_symbols = set(source["required_anchor_symbols"])
    cross_rows = []
    for timestamp, group in panel.groupby("timestamp", sort=True, observed=True):
        available = group.loc[group["archive_available"]]
        available_symbols = set(available["symbol"].astype(str))
        anchors_available = anchor_symbols.issubset(available_symbols)
        cross_rows.append(
            {
                "timestamp": int(timestamp),
                "available_symbol_count": int(len(available_symbols)),
                "anchors_available": anchors_available,
                "cross_sell_liquidation_usd": float(available["sell_liquidation_usd"].sum()),
                "cross_buy_liquidation_usd": float(available["buy_liquidation_usd"].sum()),
            }
        )
    cross = pd.DataFrame(cross_rows)
    cross["cross_liquidation_total_usd"] = cross["cross_sell_liquidation_usd"] + cross["cross_buy_liquidation_usd"]
    cross["usable_cross_market"] = (
        (cross["available_symbol_count"] >= int(source["minimum_available_symbols_per_cross_market_bar"]))
        & cross["anchors_available"]
    )
    panel = panel.merge(cross, on="timestamp", how="left", validate="many_to_one")
    return panel.reset_index(drop=True), cross, pd.DataFrame(quality_rows)


def build(contract_path: Path, output_dir: Path, workers: int) -> dict[str, Any]:
    contract = load_v7_contract(contract_path, root=ROOT)
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = output_dir / "manifest.json"
    if manifest_path.is_file():
        existing = json.loads(manifest_path.read_text(encoding="utf-8"))
        if existing.get("contract_sha256") == contract.sha256:
            return existing
        raise ValueError("stale V7 dataset manifest")
    raw_dir = output_dir / "raw"
    raw_dir.mkdir(parents=True, exist_ok=True)
    listing = _list_source(contract)
    listing_path = output_dir / "source_listing.parquet"
    listing.to_parquet(listing_path, index=False)
    catalog_path = output_dir / "raw_catalog.parquet"
    reusable: dict[tuple[str, str], dict[str, Any]] = {}
    if catalog_path.is_file():
        previous = pd.read_parquet(catalog_path)
        for row in previous.loc[previous["downloaded"] == True].to_dict(orient="records"):  # noqa: E712
            path = raw_dir / str(row["local_path"])
            if path.is_file() and sha256_file(path) == row.get("official_sha256"):
                reusable[(str(row["exchange_symbol"]), str(row["date"]))] = row
    rows: list[dict[str, Any]] = list(reusable.values())
    failures: list[dict[str, Any]] = []
    pending = [
        item for item in listing.to_dict(orient="records")
        if (str(item["exchange_symbol"]), str(item["date"])) not in reusable
    ]
    with ThreadPoolExecutor(max_workers=max(1, int(workers))) as pool:
        futures = {pool.submit(_download, item, raw_dir): item for item in pending}
        for index, future in enumerate(as_completed(futures), start=1):
            item = futures[future]
            try:
                rows.append(future.result())
            except Exception as exc:
                failures.append({**item, "downloaded": False, "error": str(exc)})
            if index % 100 == 0 or index == len(pending):
                print(json.dumps({"downloaded": len(rows), "failed": len(failures), "listed": len(listing)}), flush=True)
    catalog = pd.DataFrame([*rows, *failures]).sort_values(["exchange_symbol", "date"], kind="stable")
    catalog.to_parquet(catalog_path, index=False)
    panel, cross, coverage = _build_panels(contract, catalog, raw_dir)
    panel_path = output_dir / "liquidation_features_4h.parquet"
    cross_path = output_dir / "cross_market_liquidations_4h.parquet"
    panel.to_parquet(panel_path, index=False)
    cross.to_parquet(cross_path, index=False)
    archive_path = output_dir / "raw_liquidations.zip"
    pack_raw_directory(raw_dir, archive_path)
    shutil.rmtree(raw_dir)
    workbook_path = output_dir / "dataset_catalog.xlsx"
    _workbook(workbook_path, {"Raw Files": catalog, "Coverage": coverage, "Cross Coverage": cross})
    artifacts = {
        "feature_panel": panel_path,
        "event_bars": cross_path,
        "raw_catalog": catalog_path,
        "source_listing": listing_path,
        "raw_archive": archive_path,
        "workbook": workbook_path,
    }
    artifact_rows = {
        key: {"path": str(path.relative_to(ROOT)), "sha256": sha256_file(path)} for key, path in artifacts.items()
    }
    dataset_sha = hashlib.sha256(
        (contract.sha256 + "".join(artifact_rows[key]["sha256"] for key in sorted(artifact_rows))).encode()
    ).hexdigest()
    manifest = {
        "schema_version": "exia.genetic_primus_liquidation_dataset.v7",
        "generated_at": datetime.now(UTC).isoformat(),
        "status": "V7_LIQUIDATION_DATASET_BUILT" if not failures else "V7_LIQUIDATION_DATASET_INCOMPLETE",
        "contract_sha256": contract.sha256,
        "source_id": contract.liquidation_source["source_id"],
        "source_family": contract.liquidation_source["source_family"],
        "listed_raw_files": int(len(listing)),
        "downloaded_raw_files": int(len(rows)),
        "failed_raw_files": int(len(failures)),
        "official_file_coverage": len(rows) / len(listing),
        "feature_rows": int(len(panel)),
        "cross_market_bars": int(len(cross)),
        "usable_cross_market_bar_coverage": float(cross["usable_cross_market"].mean()),
        "coverage": coverage.to_dict(orient="records"),
        "failures": failures,
        "dataset_sha256": dataset_sha,
        "artifacts": artifact_rows,
        "point_in_time": True,
        "official_checksums_verified": not failures,
        "missing_daily_archives_zero_imputed": False,
        "raw_archive_retained": True,
        "validation_opened": False,
        "oos_opened": False,
        "paper_allowed": False,
        "live_allowed": False,
        "orders_enabled": False,
        "promotion_authority": False,
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description="Build Genetic_Primus V7 COIN-M liquidation dataset")
    parser.add_argument("--contract", type=Path, default=DEFAULT_CONTRACT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--workers", type=int, default=24)
    args = parser.parse_args()
    result = build(args.contract.resolve(), args.output_dir.resolve(), args.workers)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
