from __future__ import annotations

import json
import math
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


WORKSPACE = Path(__file__).resolve().parents[1]
RESULTS = WORKSPACE / "Results"
REPORTS = WORKSPACE / "Reports"
ASSETS = REPORTS / "panteon_flash_assets"
DOCX_PATH = REPORTS / "Panteon_Flash_operational_report_2026-05-21.docx"


@dataclass
class RunMetrics:
    key: str
    title: str
    path: Path
    years: str
    bars: int
    pnl_pct: float
    pnl_usd: float
    max_dd_pct: float
    closed_trades: int
    selected_signals: int
    filled_signals: int
    valid_for_decision: bool = True
    note: str = ""


KEY_TITLES = {
    "PanteonFlashRetrodate": "Стартовый Flash",
    "PanteonFlashRetrodateGated": "Жесткий gate",
    "PanteonFlashRetrodateShadow5": "Shadow min 5",
    "PanteonFlashRetrodateShadow50": "Shadow min 50",
    "PanteonFlashRetrodateShadow50Cap1": "Shadow50 cap1",
    "PanteonFlashRetrodateShadow50Cap1Score05": "Shadow50 cap1 score>=0.5",
    "PanteonFlashRetrodateShadow50Cap1Score05Overext2": "Защита от overextension",
    "PanteonFlashRetrodateShadow50Cap1Score05Attribution": "Базовая атрибуция",
    "PanteonFlashRetrodateShadow50Cap1Score05Deny3": "Запрет 3 строк",
    "PanteonFlashRetrodateShadow50Cap1Score05DenyWorstRows": "Запрет худших строк",
    "PanteonFlashRetrodateShadow50Cap1Score05OnlineDegrade": "Online-degradation",
    "PanteonFlashRetrodateShadow50Cap1Score05OnlineDegradeReserveCap": "Reserve-cap baseline",
    "PanteonFlashRetrodateShadow50Cap1Score05ExperimentalActors": "Эксперимент с real-allow",
    "PanteonFlashRetrodateShadow50Cap1Score05ExperimentalShadowOnly": "Shadow-only до registry split",
    "PanteonFlashRetrodateShadow50Cap1Score05ExperimentalShadowOnlyRegistrySplit": "Shadow-only registry split",
    "PanteonFlashExperimentalHalfYearGate2026Smoke": "Gate 2026 H1",
}


def load_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


def run_key(path: Path) -> str:
    for part in path.parts:
        if part.startswith("PanteonFlash"):
            return part
    return path.parent.name


def collect_runs() -> list[RunMetrics]:
    runs: list[RunMetrics] = []
    for summary_path in RESULTS.rglob("run_summary.json"):
        key = run_key(summary_path)
        if not key.startswith("PanteonFlash"):
            continue
        out = summary_path.parent
        summary = load_json(summary_path)
        status = load_json(out / "status.json")
        live = status.get("live_session", {}) if isinstance(status, dict) else {}
        attr = load_json(out / "flash_attribution_summary.json").get("summary", {})
        years = ",".join(str(year) for year in summary.get("executed_years", []))
        note = ""
        valid = True
        if key.endswith("ExperimentalShadowOnly"):
            valid = False
            note = "Только диагностика: получено до разделения real/shadow registry."
        if "Smoke" in key or years == "2026":
            valid = False
            note = note or "Smoke / короткое диагностическое окно."
        runs.append(
            RunMetrics(
                key=key,
                title=KEY_TITLES.get(key, key.replace("PanteonFlash", "")),
                path=out,
                years=years,
                bars=int(summary.get("bars_processed", 0) or 0),
                pnl_pct=float(live.get("panteon_owned_pnl_pct") or 0.0),
                pnl_usd=float(live.get("panteon_owned_realized_pnl_usd") or 0.0),
                max_dd_pct=float(status.get("panteon_max_drawdown_pct") or 0.0),
                closed_trades=int(live.get("real_closed_trades", 0) or 0),
                selected_signals=int(attr.get("selected_signals", 0) or 0),
                filled_signals=int(attr.get("filled_signals", 0) or 0),
                valid_for_decision=valid,
                note=note,
            )
        )
    return sorted(runs, key=lambda item: item.path.stat().st_mtime)


def latest_run(runs: list[RunMetrics], key: str) -> RunMetrics | None:
    matches = [run for run in runs if run.key == key]
    return matches[-1] if matches else None


def actor_summary(run: RunMetrics | None) -> list[dict[str, Any]]:
    if run is None:
        return []
    data = load_json(run.path / "flash_attribution_summary.json")
    rows = []
    for key, row in (data.get("actor_summary", {}) or {}).items():
        if not isinstance(row, dict):
            continue
        rows.append(
            {
                "key": key,
                "label": row.get("actor_label") or key,
                "type": row.get("actor_type") or "",
                "selected": int(row.get("selected_signals", 0) or 0),
                "filled": int(row.get("filled_signals", 0) or 0),
                "closed": int(row.get("closed_trades", 0) or 0),
                "pnl": float(row.get("realized_pnl_usd", 0.0) or 0.0),
            }
        )
    return sorted(rows, key=lambda row: row["pnl"], reverse=True)


def gate_report(run: RunMetrics | None) -> dict[str, Any]:
    if run is None:
        return {}
    return load_json(run.path / "experimental_flash_shadow_report.json")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def draw_text(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, *, size: int, fill: str = "#111827", bold: bool = False) -> None:
    draw.text(xy, text, fill=fill, font=font(size, bold=bold))


def sanitize_label(text: str, max_len: int = 36) -> str:
    return text if len(text) <= max_len else text[: max_len - 1] + "..."


def bar_chart(
    items: list[tuple[str, float]],
    path: Path,
    *,
    title: str,
    subtitle: str = "",
    unit: str = "%",
    width: int = 1500,
    height: int = 900,
) -> None:
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw_text(draw, (50, 35), title, size=34, bold=True, fill="#0B2545")
    if subtitle:
        draw_text(draw, (50, 78), subtitle, size=20, fill="#4B5563")
    if not items:
        draw_text(draw, (50, 160), "No data", size=24)
        image.save(path)
        return
    left = 430
    right = width - 120
    top = 145
    row_h = max(34, min(58, (height - top - 70) // max(1, len(items))))
    values = [value for _, value in items]
    min_v = min(0.0, min(values))
    max_v = max(0.0, max(values))
    span = max(1e-9, max_v - min_v)
    zero_x = int(left + (0 - min_v) / span * (right - left))
    draw.line((zero_x, top - 15, zero_x, height - 45), fill="#9CA3AF", width=2)
    for idx, (label, value) in enumerate(items):
        y = top + idx * row_h
        draw_text(draw, (50, y + 4), sanitize_label(label), size=18, fill="#111827")
        x = int(left + (value - min_v) / span * (right - left))
        color = "#1F7A4D" if value >= 0 else "#B42318"
        x1, x2 = sorted((zero_x, x))
        draw.rectangle((x1, y + 7, x2, y + row_h - 8), fill=color)
        value_text = f"{value:+.2f}{unit}"
        tx = min(right - 95, max(left + 8, x + (8 if value >= 0 else -100)))
        draw_text(draw, (tx, y + 5), value_text, size=17, fill="#111827", bold=True)
    image.save(path, quality=95)


def paired_bar_chart(
    items: list[tuple[str, float, float]],
    path: Path,
    *,
    title: str,
    left_name: str,
    right_name: str,
    width: int = 1500,
    height: int = 900,
) -> None:
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw_text(draw, (50, 35), title, size=34, bold=True, fill="#0B2545")
    draw.rectangle((50, 88, 80, 110), fill="#1F7A4D")
    draw_text(draw, (90, 84), left_name, size=18)
    draw.rectangle((270, 88, 300, 110), fill="#2E74B5")
    draw_text(draw, (310, 84), right_name, size=18)
    left = 430
    right = width - 130
    top = 145
    row_h = max(42, min(64, (height - top - 70) // max(1, len(items))))
    min_v = min(0.0, min(min(a, b) for _, a, b in items))
    max_v = max(1.0, max(max(a, b) for _, a, b in items))
    span = max(1.0, max_v - min_v)
    zero_x = int(left + (0 - min_v) / span * (right - left))
    draw.line((zero_x, top - 15, zero_x, height - 45), fill="#9CA3AF", width=2)
    for idx, (label, a, b) in enumerate(items):
        y = top + idx * row_h
        draw_text(draw, (50, y + 9), sanitize_label(label), size=18)
        ax = int(left + (a - min_v) / span * (right - left))
        bx = int(left + (b - min_v) / span * (right - left))
        a1, a2 = sorted((zero_x, ax))
        b1, b2 = sorted((zero_x, bx))
        draw.rectangle((a1, y + 6, a2, y + 23), fill="#1F7A4D" if a >= 0 else "#B42318")
        draw.rectangle((b1, y + 28, b2, y + 45), fill="#2E74B5" if b >= 0 else "#B42318")
        atx = min(right - 60, max(left + 8, ax + (8 if a >= 0 else -65)))
        btx = min(right - 60, max(left + 8, bx + (8 if b >= 0 else -65)))
        draw_text(draw, (atx, y + 3), f"{a:.1f}", size=15)
        draw_text(draw, (btx, y + 25), f"{b:.1f}", size=15)
    image.save(path, quality=95)


def scatter_chart(items: list[RunMetrics], path: Path) -> None:
    width, height = 1400, 850
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw_text(draw, (55, 35), "Доходность vs просадка, 5-летние прогоны", size=34, bold=True, fill="#0B2545")
    plot = (140, 130, width - 90, height - 120)
    draw.rectangle(plot, outline="#CBD5E1", width=2)
    max_dd = max(35.0, max((run.max_dd_pct for run in items), default=0.0))
    min_pnl = min(-40.0, min((run.pnl_pct for run in items), default=0.0))
    max_pnl = max(30.0, max((run.pnl_pct for run in items), default=0.0))
    def xy(run: RunMetrics) -> tuple[int, int]:
        x = plot[0] + int(run.max_dd_pct / max_dd * (plot[2] - plot[0]))
        y = plot[3] - int((run.pnl_pct - min_pnl) / (max_pnl - min_pnl) * (plot[3] - plot[1]))
        return x, y
    zero_y = plot[3] - int((0 - min_pnl) / (max_pnl - min_pnl) * (plot[3] - plot[1]))
    draw.line((plot[0], zero_y, plot[2], zero_y), fill="#9CA3AF", width=2)
    combined_points: dict[str, tuple[int, int]] = {}
    for run in items:
        x, y = xy(run)
        color = "#1F7A4D" if run.pnl_pct > 0 and run.valid_for_decision else "#B42318"
        if not run.valid_for_decision:
            color = "#7A5A00"
        radius = 7 + min(16, int(math.sqrt(max(0, run.closed_trades)) / 3))
        draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=color, outline="#111827")
        if run.key == "PanteonFlashRetrodateShadow50Cap1Score05OnlineDegradeReserveCap":
            combined_points["best"] = (x, y)
        elif run.key == "PanteonFlashRetrodateShadow50Cap1Score05ExperimentalShadowOnlyRegistrySplit":
            combined_points["split"] = (x, y)
        elif run.key == "PanteonFlashRetrodateShadow50Cap1Score05ExperimentalActors":
            draw_text(draw, (x + 12, y - 12), sanitize_label(run.title, 28), size=15, fill="#111827")
    if "best" in combined_points:
        x, y = combined_points["best"]
        label = "Reserve-cap baseline / registry split" if "split" in combined_points else "Reserve-cap baseline"
        draw_text(draw, (x + 12, y - 12), sanitize_label(label, 38), size=15, fill="#111827")
    draw_text(draw, (plot[0], height - 78), "X: максимальная просадка, ниже лучше", size=18, fill="#4B5563")
    draw_text(draw, (plot[0], height - 52), "Y: PnL %, выше лучше. Золото = только диагностика/smoke.", size=18, fill="#4B5563")
    image.save(path, quality=95)


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color in [
        ("Title", 24, "0B2545"),
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Heading 3" or True
    return doc


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "D0D7DE")
        borders.append(tag)
    tbl_pr.append(borders)


def add_table(doc: Document, headers: list[str], rows: list[list[Any]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0].cells
    for idx, text in enumerate(headers):
        header[idx].text = str(text)
        set_cell_shading(header[idx], "F2F4F7")
        for p in header[idx].paragraphs:
            p.runs[0].font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx > 0:
                for p in cells[idx].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_image(doc: Document, path: Path, caption: str) -> None:
    doc.add_picture(str(path), width=Inches(6.3))
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.italic = True
        run.font.color.rgb = RGBColor.from_string("4B5563")


def fmt_pct(value: float) -> str:
    return f"{value:+.2f}%"


def fmt_num(value: float) -> str:
    return f"{value:.2f}"


def create_charts(runs: list[RunMetrics], best: RunMetrics, exp: RunMetrics | None, smoke: RunMetrics | None) -> dict[str, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    full_runs = [
        run
        for run in runs
        if run.years == "2022,2023,2024,2025,2026"
        and run.key.startswith("PanteonFlash")
    ]
    full_for_bar = [
        (run.title, run.pnl_pct)
        for run in full_runs
        if run.key not in {"PanteonFlashRetrodateShadow50Cap1Score05ExperimentalShadowOnly"}
    ]
    charts = {}
    charts["timeline"] = ASSETS / "full_run_pnl_timeline.png"
    bar_chart(
        full_for_bar,
        charts["timeline"],
        title="5-летний PnL по экспериментам",
        subtitle="Полные прогоны 2022-2026 с часовым шагом; невалидный pre-split прогон исключен",
    )
    charts["scatter"] = ASSETS / "pnl_vs_drawdown.png"
    scatter_chart(full_runs, charts["scatter"])
    best_actors = actor_summary(best)
    charts["best_actor_pnl"] = ASSETS / "best_actor_pnl.png"
    bar_chart(
        [(row["label"], row["pnl"]) for row in best_actors],
        charts["best_actor_pnl"],
        title="Лучший baseline: реализованный PnL по акторам",
        subtitle=best.title,
        unit="$",
    )
    if exp:
        charts["exp_actor_pnl"] = ASSETS / "experimental_actor_pnl.png"
        bar_chart(
            [(row["label"], row["pnl"]) for row in actor_summary(exp)],
            charts["exp_actor_pnl"],
            title="Эксперимент real-allow: вклад акторов",
            subtitle="Текущие wrapper-агенты изменили real allocation и ухудшили результат",
            unit="$",
        )
    if smoke:
        gate = gate_report(smoke)
        rows = gate.get("labels", []) if isinstance(gate, dict) else []
        charts["gate"] = ASSETS / "experimental_gate_2026.png"
        paired_bar_chart(
            [
                (str(row.get("label")), float(row.get("shadow_pnl_pct", 0.0)), float(row.get("max_drawdown_pct", 0.0)))
                for row in rows
            ],
            charts["gate"],
            title="2026 H1 shadow gate: wrapper-агенты",
            left_name="PnL %",
            right_name="Max DD %",
        )
    return charts


def build_report() -> Path:
    REPORTS.mkdir(exist_ok=True)
    runs = collect_runs()
    best = latest_run(runs, "PanteonFlashRetrodateShadow50Cap1Score05OnlineDegradeReserveCap")
    exp = latest_run(runs, "PanteonFlashRetrodateShadow50Cap1Score05ExperimentalActors")
    split = latest_run(runs, "PanteonFlashRetrodateShadow50Cap1Score05ExperimentalShadowOnlyRegistrySplit")
    smoke = latest_run(runs, "PanteonFlashExperimentalHalfYearGate2026Smoke")
    if best is None:
        raise SystemExit("Best baseline run was not found")
    charts = create_charts(runs, best, exp, smoke)

    doc = setup_doc()
    title = doc.add_paragraph("Panteon Flash: отчет о готовности к эксплуатации", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph(
        "Анализ ретропрогонов 2022-2026, экспериментальных агентов/игроков и готовности к реальным торгам. "
        f"Сформировано {datetime.now().strftime('%Y-%m-%d %H:%M')} по локальному времени."
    )
    subtitle.runs[0].font.color.rgb = RGBColor.from_string("4B5563")

    doc.add_heading("Ключевое решение", level=1)
    p = doc.add_paragraph()
    run = p.add_run("Решение: полноценные реальные торги пока не запускать.")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("9B1C1C")
    doc.add_paragraph(
        "Лучшая пригодная к деплою конфигурация Panteon_Flash прибыльна на 5-летнем replay, "
        "но текущие экспериментальные wrapper-агенты не прошли последний полугодовой gate, а real-allow эксперимент дал крупную просадку. "
        "Следующий безопасный шаг - paper trading или строго ограниченный canary только после включения эксплуатационных контролей."
    )
    add_bullets(
        doc,
        [
            f"Лучший валидный 5-летний прогон: {best.title}, PnL {fmt_pct(best.pnl_pct)}, max DD {best.max_dd_pct:.2f}%, закрытых сделок {best.closed_trades}.",
            "Экспериментальные wrapper-агенты в real-allow не готовы: 5-летний тест ушел в -11.69% при max DD 29.05%.",
            "Режим shadow-only registry split воспроизвел лучший baseline и не дает диагностическим агентам попадать в real allocation.",
            "Gate 2026 H1 не продвинул ни одного экспериментального wrapper-а, поэтому текущие варианты должны оставаться shadow-only.",
        ],
    )

    doc.add_heading("Ключевые результаты", level=1)
    add_table(
        doc,
        ["Прогон", "Годы", "PnL", "Max DD", "Закрыто", "Выбрано", "Роль в решении"],
        [
            [
                best.title,
                best.years,
                fmt_pct(best.pnl_pct),
                f"{best.max_dd_pct:.2f}%",
                best.closed_trades,
                best.selected_signals,
                "Кандидат baseline",
            ],
            [
                split.title if split else "Registry split",
                split.years if split else "-",
                fmt_pct(split.pnl_pct) if split else "-",
                f"{split.max_dd_pct:.2f}%" if split else "-",
                split.closed_trades if split else "-",
                split.selected_signals if split else "-",
                "Проверка shadow-only изоляции",
            ],
            [
                exp.title if exp else "Эксперимент с real-allow",
                exp.years if exp else "-",
                fmt_pct(exp.pnl_pct) if exp else "-",
                f"{exp.max_dd_pct:.2f}%" if exp else "-",
                exp.closed_trades if exp else "-",
                exp.selected_signals if exp else "-",
                "Отклонить",
            ],
            [
                smoke.title if smoke else "Gate 2026 smoke",
                smoke.years if smoke else "-",
                fmt_pct(smoke.pnl_pct) if smoke else "-",
                f"{smoke.max_dd_pct:.2f}%" if smoke else "-",
                smoke.closed_trades if smoke else "-",
                smoke.selected_signals if smoke else "-",
                "Диагностика early-stop",
            ],
        ],
    )
    add_image(doc, charts["timeline"], "Рисунок 1. Полные 5-летние эксперименты по реализованному PnL.")
    add_image(doc, charts["scatter"], "Рисунок 2. Доходность и просадка для 5-летних экспериментов.")

    doc.add_heading("Атрибуция лучшего baseline", level=1)
    doc.add_paragraph(
        "Вариант с reserve-cap degradation - лучший пригодный к деплою прогон. Он сохраняет упрощенный путь Panteon_Flash и ограничивает пере-концентрацию, когда degraded-актор все еще имеет сильный резервный сигнал."
    )
    best_actors = actor_summary(best)
    add_table(
        doc,
        ["Актор", "Выбрано", "Исполнено", "Закрыто", "Реализованный PnL"],
        [
            [row["label"], row["selected"], row["filled"], row["closed"], f"${row['pnl']:.2f}"]
            for row in best_actors
        ],
    )
    add_image(doc, charts["best_actor_pnl"], "Рисунок 3. Вклад акторов в реализованный PnL лучшего baseline.")

    doc.add_heading("Экспериментальные агенты и игроки", level=1)
    doc.add_paragraph(
        "Экспериментальные wrapper-агенты полезны для диагностики, но их нельзя включать в real allocation. "
        "Когда им разрешили торговать, они изменили смесь акторов и проявили токсичный spot-buy путь MomentumScalper. "
        "После registry split они остаются видимыми в shadow-метриках, но не влияют на real-кандидатов."
    )
    if exp:
        add_image(doc, charts["exp_actor_pnl"], "Рисунок 4. Эксперимент real-allow: вклад акторов.")
    if smoke:
        gate = gate_report(smoke)
        gate_rows = gate.get("labels", []) if isinstance(gate, dict) else []
        add_table(
            doc,
            ["Wrapper", "PnL 2026 H1", "Max DD", "Закрыто", "Gate", "Причина"],
            [
                [
                    row.get("label", "-"),
                    f"{float(row.get('shadow_pnl_pct', 0.0)):.2f}%",
                    f"{float(row.get('max_drawdown_pct', 0.0)):.2f}%",
                    int(row.get("closed_trades", 0) or 0),
                    "pass" if row.get("shadow_gate_passed") else "fail",
                    ", ".join(row.get("gate_reasons", []) or []),
                ]
                for row in gate_rows
            ],
        )
        add_image(doc, charts["gate"], "Рисунок 5. Gate экспериментальных wrapper-агентов за 2026 H1.")

    doc.add_heading("Можно ли начинать реальные торги?", level=1)
    add_table(
        doc,
        ["Критерий", "Статус", "Доказательство / пробел"],
        [
            ["5-летний edge в симуляции", "Частичный pass", f"Лучший валидный прогон: {fmt_pct(best.pnl_pct)}, DD {best.max_dd_pct:.2f}%."],
            ["Продвижение экспериментальных агентов", "Fail", "Gate 2026 H1 не продвинул ни одного wrapper-а."],
            ["Изоляция диагностики", "Pass", "Registry split не допускает shadow-only агентов в real selection."],
            ["Операционный мониторинг", "Не завершено", "Нужны live alerts, репетиция kill-switch, сверка позиций с биржей."],
            ["Forward validation", "Не завершено", "Нужен период paper/live-sim после code freeze."],
            ["План риска по капиталу", "Не завершено", "Нужны лимиты canary, max daily loss, rollback plan и ручная процедура stop."],
        ],
    )
    doc.add_paragraph(
        "Инженерная рекомендация: не запускать неограниченные реальные торги. Начать с paper trading на registry-split сборке, затем переходить к жестко ограниченному canary только если live paper-метрики совпадают с replay-gate."
    )

    doc.add_heading("Необходимые следующие шаги", level=1)
    add_bullets(
        doc,
        [
            "Заморозить текущую лучшую baseline-конфигурацию как кандидата для деплоя: shadow confirmation, максимум один сигнал на актора, online degradation guard и reserve actor cap.",
            "Запустить paper-trading soak с той же конфигурацией Panteon_Flash и архивировать ежедневные gate-отчеты.",
            "Добавить live operations checks: обнаружение desync с биржей, stale-feed alarm, API error alarm, max daily loss kill switch и runbook ручной аварийной остановки.",
            "Создать promotion manifest: в real allocation могут попадать только actor/symbol/action ключи, прошедшие rolling shadow gates.",
            "Держать все экспериментальные wrapper-агенты в shadow-only до прохождения latest-period и multi-period gates.",
            "Перед любым canary требовать clean startup, clean shutdown, paper trading PnL/DD gate, отсутствие desync и детерминированную воспроизводимость replay.",
        ],
    )

    doc.add_heading("Как улучшить Пантеон, агентов и игроков", level=1)
    doc.add_heading("Panteon", level=2)
    add_bullets(
        doc,
        [
            "Продвигать не целых агентов, а actor/symbol/action ключи. Атрибуция показывает, что один агент может содержать и сильные, и токсичные под-поведения.",
            "Сохранить простую архитектуру: один лучший актор на symbol/action, но добавить promotion manifest и rolling gate вокруг candidate set.",
            "Использовать latest-period gates как первое условие остановки до полного 5-летнего replay.",
            "Добавить shadow attribution по symbol/action, чтобы будущая диагностика wrapper-ов не выводила токсичность только из real attribution.",
        ],
    )
    doc.add_heading("Агенты", level=2)
    add_bullets(
        doc,
        [
            "Не продвигать текущие MomentumScalper wrapper-ы: они провалили 2026 H1.",
            "Отдельно исследовать токсичность MomentumScalper spot-buy и short-поведение. Лучший baseline в целом заработал, но spot-buy строки нестабильны по символам и периодам.",
            "ResearchValidatorAgent показывает сильный shadow-результат, но mixed real attribution; перед продвижением нужен более строгий executable gating.",
            "VolBreakoutHunter силен в shadow, но VolBreakoutSpotOnly провалил последний gate; нельзя делать простое action slicing без периодной валидации.",
        ],
    )
    doc.add_heading("Игроки", level=2)
    add_bullets(
        doc,
        [
            "Оставить игроков ансамблевыми сущностями для отдельной статистики и визуализаций, а не как дополнительный слой выбора.",
            "Использовать fixed/rotating экспериментальных игроков только в shadow или явных what-if прогонах до прохождения latest-period gates.",
            "Добавить player-level полугодовые gates на той же базе, что и gates для экспериментальных агентов.",
        ],
    )

    doc.add_heading("Как расширить тесты", level=1)
    add_bullets(
        doc,
        [
            "Добавить replay smoke gates для tiers 6 месяцев, 1 год и 5 лет; автоматически останавливаться, если latest-period gate провален.",
            "Добавить regression tests для promotion manifest: shadow-only labels не должны попадать в real attribution без явного разрешения.",
            "Добавить symbol/action attribution tests для shadow tournament events после добавления symbol/action в shadow PnL streams.",
            "Добавить exchange fault tests: stale feed, отказ ордера, partial fill, серия API errors и восстановление после position desync.",
            "Добавить deterministic replay checks: одинаковые seed/config/data должны воспроизводить PnL, число selected signals и actor attribution.",
            "Добавить risk tests вокруг max daily loss, max new opens per bar и сохранения kill-switch после restart.",
        ],
    )

    doc.add_heading("Приложение: исходные артефакты", level=1)
    add_bullets(
        doc,
        [
            f"Лучший baseline: {best.path}",
            f"Проверка registry split: {split.path if split else '-'}",
            f"Отклоненный real-allow эксперимент: {exp.path if exp else '-'}",
            f"Последний half-year gate smoke 2026: {smoke.path if smoke else '-'}",
            "Все метрики прочитаны из локальных артефактов Results: run_summary/status/attribution/gate JSON.",
        ],
    )

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    path = build_report()
    print(path)
