from __future__ import annotations

import argparse
import importlib.util
import json
from pathlib import Path
from typing import Any


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Reports" / "PanteonFlashFinalTradingReadiness_20260523"
CHARTS = OUT / "charts"


EXPERIMENTS = [
    (
        "2025",
        "base_cap10",
        "Base cap10",
        "Fail",
        ROOT / r"Results\PanteonFlashGeneratedSelectedDeny2025_20260522\proven_solo_portfolio_hard_solo_momentum_cap10\RETRODATE_MARKET\2026-05-22_19-38-36_retrodate_market_v2",
        "Базовый Flash выбирал плохой selected subset: общая архитектура была прибыльной в компонентах, но выбранные сделки не переносили edge.",
    ),
    (
        "2025",
        "generated_selected_deny",
        "Selected deny v1",
        "Intermediate",
        ROOT / r"Results\PanteonFlashGeneratedSelectedDeny2025_20260522\proven_solo_portfolio_hard_solo_momentum_cap10_generated_selected_deny\RETRODATE_MARKET\2026-05-22_19-50-35_retrodate_market_v2",
        "Первый слой запрета худших selected signal_key резко улучшил результат, но Пантеон всё ещё уступал лучшему компоненту.",
    ),
    (
        "2025",
        "round2",
        "Round2 deny",
        "Success",
        ROOT / r"Results\PanteonFlashGeneratedSelectedDenyRound2_2025_20260522\RETRODATE_MARKET\2026-05-22_20-05-53_retrodate_market_v2",
        "Второй слой selected-deny впервые дал доминацию Пантеона над лучшей составляющей на 2025.",
    ),
    (
        "2025",
        "shadow_round2only_fail",
        "Broad shadow deny",
        "Fail",
        ROOT / r"Results\PanteonFlashGeneratedSelectedDenyShadowFullRound2Only_2025_20260523\proven_solo_portfolio_hard_solo_momentum_cap10_generated_selected_deny_shadow_full_round2only\RETRODATE_MARKET\2026-05-22_21-38-38_retrodate_market_v2",
        "Широкий full-shadow deny перерезал полезные ключи и вызвал over-selection fallback-кандидатов.",
    ),
    (
        "2025",
        "atom_deny_wrong_base",
        "ATOM deny without round2",
        "Fail",
        ROOT / r"Results\PanteonFlashSoloLiveCrashAtomDeny_2025_20260523\proven_solo_portfolio_hard_solo_momentum_cap10_deny_solo_livecrash_atom_short\RETRODATE_MARKET\2026-05-22_21-56-24_retrodate_market_v2",
        "Одиночный ATOM deny без round2-базы не заменяет уже обученный слой запретов и ломает 2025.",
    ),
    (
        "2025",
        "round2_plus_terminal_atom",
        "Round2 + terminal ATOM",
        "Recommended",
        ROOT / r"Results\PanteonFlashRound2PlusTerminalAtom_2025_20260523\proven_solo_portfolio_hard_solo_momentum_cap10_round2_plus_terminal_atom\RETRODATE_MARKET\2026-05-22_22-35-42_retrodate_market_v2",
        "Финальный устойчивый кандидат: сохраняет 2025 dominance и добавляет узкий guard против OOS-проблемы ATOM.",
    ),
    (
        "2026 H1",
        "round2",
        "Round2 deny",
        "Intermediate",
        ROOT / r"Results\PanteonFlashGeneratedSelectedDenyRound2_2026H1_20260523\RETRODATE_MARKET\2026-05-22_21-02-19_retrodate_market_v2",
        "OOS прибыльный, но strict alpha 2 п.п. не достигнут.",
    ),
    (
        "2026 H1",
        "round2_open2_fail",
        "Round2 open2",
        "Fail",
        ROOT / r"Results\PanteonFlashGeneratedSelectedDenyRound2Open2_2026H1_20260523\RETRODATE_MARKET\2026-05-22_21-07-29_retrodate_market_v2",
        "Расширение числа новых открытий ухудшило OOS: прежние лимиты защищали от плохой экспозиции.",
    ),
    (
        "2026 H1",
        "shadow_3sources_fail",
        "Shadow deny 3 sources",
        "Fail",
        ROOT / r"Results\PanteonFlashGeneratedSelectedDenyShadowFull_2026H1_20260523\proven_solo_portfolio_hard_solo_momentum_cap10_generated_selected_deny_shadow_full\RETRODATE_MARKET\2026-05-22_21-29-03_retrodate_market_v2",
        "Слишком широкий deny из нескольких источников ухудшил результат относительно round2.",
    ),
    (
        "2026 H1",
        "shadow_round2only_fail_is2025",
        "Shadow deny round2-only",
        "Overfit",
        ROOT / r"Results\PanteonFlashGeneratedSelectedDenyShadowFullRound2Only_2026H1_20260523\proven_solo_portfolio_hard_solo_momentum_cap10_generated_selected_deny_shadow_full_round2only\RETRODATE_MARKET\2026-05-22_21-33-49_retrodate_market_v2",
        "На 2026 H1 выглядит лучше финала, но провален 2025-контроль, поэтому непригоден как production policy.",
    ),
    (
        "2026 H1",
        "atom_deny_wrong_base",
        "ATOM deny without round2",
        "Overfit",
        ROOT / r"Results\PanteonFlashSoloLiveCrashAtomDeny_2026H1_20260523\proven_solo_portfolio_hard_solo_momentum_cap10_deny_solo_livecrash_atom_short\RETRODATE_MARKET\2026-05-22_21-52-07_retrodate_market_v2",
        "Лучший одиночный OOS, но 2025 провален; это пример подгонки под одну проблему.",
    ),
    (
        "2026 H1",
        "round2_plus_terminal_atom",
        "Round2 + terminal ATOM",
        "Recommended",
        ROOT / r"Results\PanteonFlashRound2PlusTerminalAtom_2026H1_20260523\proven_solo_portfolio_hard_solo_momentum_cap10_round2_plus_terminal_atom\RETRODATE_MARKET\2026-05-22_22-31-22_retrodate_market_v2",
        "Лучший устойчивый вариант: проходит 2025 и 2026 H1, strict alpha и churn budget.",
    ),
]


def _load_matrix_tool() -> Any:
    spec = importlib.util.spec_from_file_location(
        "panteon_flash_matrix_tool",
        ROOT / "tools" / "run_panteon_flash_profitability_matrix.py",
    )
    if spec is None or spec.loader is None:
        raise RuntimeError("Cannot load profitability matrix tool")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _rows() -> list[dict[str, Any]]:
    module = _load_matrix_tool()
    rows: list[dict[str, Any]] = []
    for period, name, label, category, path, note in EXPERIMENTS:
        metrics = module._extract_metrics(path)
        rows.append({
            "period": period,
            "name": name,
            "label": label,
            "category": category,
            "path": str(path),
            "note": note,
            "pnl_pct": metrics.get("pnl_pct"),
            "best_component_pnl_pct": metrics.get("best_component_pnl_pct"),
            "panteon_alpha_pct": metrics.get("panteon_alpha_pct"),
            "beats_best_component": metrics.get("beats_best_component"),
            "panteon_lcb_dominance": metrics.get("panteon_lcb_dominance"),
            "ensemble_lift_pass": metrics.get("ensemble_lift_pass"),
            "regime_floor_pass": metrics.get("regime_floor_pass"),
            "regime_floor_insufficient_regimes": metrics.get("regime_floor_insufficient_regimes"),
            "churn_budget_pass": metrics.get("churn_budget_pass"),
            "actor_switches_per_day": metrics.get("actor_switches_per_day"),
            "flash_selected_signals": metrics.get("flash_selected_signals"),
            "flash_executable_selected_signals": metrics.get("flash_executable_selected_signals"),
            "dominance_equity_ratio": metrics.get("dominance_equity_ratio"),
        })
    return rows


def build_data() -> list[dict[str, Any]]:
    OUT.mkdir(parents=True, exist_ok=True)
    CHARTS.mkdir(exist_ok=True)
    rows = _rows()
    (OUT / "panteon_flash_iteration_metrics.json").write_text(
        json.dumps(rows, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    headers = [
        "period",
        "label",
        "category",
        "pnl_pct",
        "best_component_pnl_pct",
        "panteon_alpha_pct",
        "beats_best_component",
        "panteon_lcb_dominance",
        "churn_budget_pass",
        "actor_switches_per_day",
        "flash_executable_selected_signals",
        "note",
        "path",
    ]
    lines = [";".join(headers)]
    for row in rows:
        lines.append(";".join(str(row.get(header, "")).replace(";", ",") for header in headers))
    (OUT / "panteon_flash_iteration_metrics.csv").write_text(
        "\n".join(lines),
        encoding="utf-8-sig",
    )
    return rows


def build_charts(rows: list[dict[str, Any]]) -> None:
    import matplotlib.pyplot as plt

    colors = {
        "Fail": "#C0392B",
        "Intermediate": "#7F8C8D",
        "Success": "#2E86AB",
        "Overfit": "#D68910",
        "Recommended": "#1E8449",
    }
    plt.rcParams.update({"font.size": 9, "axes.titlesize": 12, "axes.labelsize": 10})

    fig, axes = plt.subplots(1, 2, figsize=(13, 5), sharey=True)
    for ax, period in zip(axes, ["2025", "2026 H1"]):
        part = [row for row in rows if row["period"] == period]
        x = list(range(len(part)))
        ax.bar(
            [i - 0.18 for i in x],
            [row["pnl_pct"] for row in part],
            width=0.36,
            label="Flash PnL",
            color=[colors[row["category"]] for row in part],
        )
        ax.bar(
            [i + 0.18 for i in x],
            [row["best_component_pnl_pct"] for row in part],
            width=0.36,
            label="Best component",
            color="#AAB7B8",
        )
        ax.axhline(0, color="#333333", linewidth=0.8)
        ax.set_title(period)
        ax.set_xticks(x)
        ax.set_xticklabels([row["label"] for row in part], rotation=35, ha="right")
        ax.set_ylabel("PnL, %")
        ax.grid(axis="y", alpha=0.25)
    axes[0].legend(loc="upper left")
    fig.suptitle("Flash vs Best Standalone Component")
    fig.tight_layout()
    fig.savefig(CHARTS / "01_pnl_vs_best_component.png", dpi=180)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(12, 5))
    display = [row["period"] + " / " + row["label"] for row in rows]
    ax.barh(display, [row["panteon_alpha_pct"] for row in rows], color=[colors[row["category"]] for row in rows])
    ax.axvline(0, color="#333333", linewidth=0.8)
    ax.axvline(2.0, color="#1E8449", linewidth=1.2, linestyle="--", label="strict alpha 2 pp")
    ax.set_xlabel("Alpha vs best component, percentage points")
    ax.set_title("Where Panteon Beats Its Components")
    ax.grid(axis="x", alpha=0.25)
    ax.legend(loc="lower right")
    fig.tight_layout()
    fig.savefig(CHARTS / "02_alpha_by_iteration.png", dpi=180)
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(13, 5))
    selected_names = {
        "generated_selected_deny",
        "round2",
        "round2_plus_terminal_atom",
        "round2_open2_fail",
    }
    subset = [row for row in rows if row["name"] in selected_names]
    labels = [row["period"] + " / " + row["label"] for row in subset]
    axes[0].barh(labels, [row["flash_executable_selected_signals"] for row in subset], color="#2E86AB")
    axes[0].set_title("Executable Selected Signals")
    axes[0].set_xlabel("Count")
    axes[0].grid(axis="x", alpha=0.25)
    axes[1].barh(labels, [row["actor_switches_per_day"] for row in subset], color="#7D3C98")
    axes[1].axvline(24.0, color="#C0392B", linestyle="--", linewidth=1.0, label="churn budget")
    axes[1].set_title("Actor Switches per Day")
    axes[1].set_xlabel("Switches/day")
    axes[1].grid(axis="x", alpha=0.25)
    fig.suptitle("Activity Control: Limits Protected the Portfolio")
    fig.tight_layout()
    fig.savefig(CHARTS / "03_signals_and_churn.png", dpi=180)
    plt.close(fig)

    gates = [
        ("beats_best_component", "Best comp"),
        ("panteon_lcb_dominance", "LCB"),
        ("ensemble_lift_pass", "Ensemble"),
        ("churn_budget_pass", "Churn"),
    ]
    gate_names = {
        "base_cap10",
        "generated_selected_deny",
        "round2",
        "round2_plus_terminal_atom",
        "shadow_round2only_fail",
        "atom_deny_wrong_base",
        "round2_open2_fail",
    }
    gate_rows = [row for row in rows if row["name"] in gate_names]
    matrix = []
    for row in gate_rows:
        values = []
        for key, _label in gates:
            value = row.get(key)
            values.append(1 if value is True else 0 if value is False else 0.5)
        matrix.append(values)
    fig, ax = plt.subplots(figsize=(10, 4.5))
    ax.imshow(matrix, aspect="auto", cmap="RdYlGn", vmin=0, vmax=1)
    ax.set_xticks(range(len(gates)))
    ax.set_xticklabels([label for _key, label in gates])
    ax.set_yticks(range(len(gate_rows)))
    ax.set_yticklabels([row["period"] + " / " + row["label"] for row in gate_rows])
    for i, values in enumerate(matrix):
        for j, value in enumerate(values):
            label = "PASS" if value == 1 else "FAIL" if value == 0 else "N/A"
            ax.text(j, i, label, ha="center", va="center", fontsize=7, color="#111111")
    ax.set_title("Production Gate Matrix")
    fig.tight_layout()
    fig.savefig(CHARTS / "04_gate_matrix.png", dpi=180)
    plt.close(fig)


def _fmt(value: Any, digits: int = 2) -> str:
    if isinstance(value, bool):
        return "да" if value else "нет"
    if value is None:
        return "n/a"
    if isinstance(value, (int, float)):
        return f"{float(value):.{digits}f}"
    return str(value)


def build_markdown(rows: list[dict[str, Any]]) -> Path:
    final_2025 = next(row for row in rows if row["period"] == "2025" and row["name"] == "round2_plus_terminal_atom")
    final_2026 = next(row for row in rows if row["period"] == "2026 H1" and row["name"] == "round2_plus_terminal_atom")
    md = OUT / "Panteon_Flash_final_trading_readiness_report.md"
    lines = [
        "# Panteon Flash: итоговый отчёт по готовности к реальным торгам",
        "",
        "Дата: 2026-05-23",
        "",
        "## Executive Summary",
        "",
        "Лучший устойчивый вариант для следующего этапа: `proven_solo_portfolio_hard_solo_momentum_cap10_generated_selected_deny_round2_terminal_atom`.",
        "Это не рекомендация запускать неограниченные реальные торги. Результат достаточен для ограниченного live/paper pilot с малыми лимитами, обязательным shadow-мониторингом и автоматическим kill-switch.",
        "",
        "Ключевой вывод: Пантеон стал прибыльнее своих составляющих только после перехода от общего выбора акторов к контролю конкретных `actor|symbol|action` ключей и после сохранения жёстких лимитов экспозиции.",
        "",
        "## Лучший результат",
        "",
        "| Период | Flash PnL | Best component | Alpha | Beats best | LCB dominance | Churn/day | Executable signals |",
        "|---|---:|---:|---:|---|---|---:|---:|",
        f"| 2025 | {_fmt(final_2025['pnl_pct'])}% | {_fmt(final_2025['best_component_pnl_pct'])}% | {_fmt(final_2025['panteon_alpha_pct'])} п.п. | {_fmt(final_2025['beats_best_component'])} | {_fmt(final_2025['panteon_lcb_dominance'])} | {_fmt(final_2025['actor_switches_per_day'])} | {final_2025['flash_executable_selected_signals']} |",
        f"| 2026 H1 | {_fmt(final_2026['pnl_pct'])}% | {_fmt(final_2026['best_component_pnl_pct'])}% | {_fmt(final_2026['panteon_alpha_pct'])} п.п. | {_fmt(final_2026['beats_best_component'])} | {_fmt(final_2026['panteon_lcb_dominance'])} | {_fmt(final_2026['actor_switches_per_day'])} | {final_2026['flash_executable_selected_signals']} |",
        "",
        "Конфигурационная логика финала:",
        "",
        "- база: `PROVEN_SOLO_PORTFOLIO_HARD_SOLO_MOMENTUM_CAP10_ARGS`;",
        "- generated selected-deny из 2025 `base_cap10` + `generated_selected_deny`;",
        "- новый terminal guard: `ensemble:Solo_LiveCrashHunter|ATOM/USDT|FUT_SHORT_FULL`;",
        "- лимит новых открытий оставлен строгим: расширение до `open2` ухудшило 2026 H1.",
        "",
        "## Визуализации",
        "",
        f"![PnL vs best component]({(CHARTS / '01_pnl_vs_best_component.png').as_posix()})",
        "",
        f"![Alpha by iteration]({(CHARTS / '02_alpha_by_iteration.png').as_posix()})",
        "",
        f"![Signals and churn]({(CHARTS / '03_signals_and_churn.png').as_posix()})",
        "",
        f"![Gate matrix]({(CHARTS / '04_gate_matrix.png').as_posix()})",
        "",
        "## Причины успехов",
        "",
        "1. `selected-deny` работал лучше, чем общий запрет акторов: проблема была не в акторе целиком, а в конкретных `actor|symbol|action` сочетаниях.",
        "2. Второй слой deny на 2025 снял selection bias базового Flash: 2025 вырос с 3.76% до 21.35%, alpha стал +3.19 п.п.",
        "3. Terminal ATOM guard устранил главный OOS-разрыв 2026 H1 без разрушения 2025, если применять его поверх round2, а не вместо round2.",
        "4. Финальный вариант прошёл оба контрольных периода: 2025 остался выше лучшего компонента, а 2026 H1 преодолел strict alpha threshold.",
        "",
        "## Причины провалов",
        "",
        "1. Broad shadow deny переобучался и ломал 2025: широкая фильтрация full-shadow статистики перерезала полезные ключи и снижала Flash PnL до 2.57%.",
        "2. Одиночный ATOM deny без round2-базы выглядел хорошо на 2026 H1, но проваливал 2025: без уже обученного round2 слоя Flash начинал выбирать плохие fallback-кандидаты.",
        "3. Расширение open-limit до open2 ухудшало OOS: 2026 H1 снижался с 10.14% до 9.54%, значит текущие лимиты реально защищают от плохой экспозиции.",
        "",
        "## Production Readiness",
        "",
        "Статус: не запускать без ограничений. Лучший вариант можно запускать только как контролируемый пилот.",
        "",
        "Минимальный режим запуска:",
        "",
        "- старт с paper/live-shadow на 1-2 недели на той же конфигурации;",
        "- затем real capital не более 5-10% от планового лимита;",
        "- `max_new_opens_per_bar=1`, не расширять до open2/open3;",
        "- kill-switch: отключить real execution при просадке Пантеона > 1.5-2.0% от equity или при `PanteonAdvantage < 0` на rolling window;",
        "- ежедневный отчёт `standalone_vs_flash_selected` и `flash_attribution_summary`;",
        "- запрещать новые generated-deny правила только после прохождения 2025 + 2026 H1 + будущего holdout.",
        "",
        "## Следующие шаги",
        "",
        "1. Прогнать финальный кандидат на полном 2022-2026 walk-forward.",
        "2. Добавить CI-gate, который запрещает merge, если 2025 или 2026 H1 не проходят `beats_best_component`, `panteon_lcb_dominance`, `churn_budget`.",
        "3. Добавить отчёт diff выбранных ключей после каждого deny/terminal-deny, чтобы видеть fallback-сделки до запуска expensive прогонов.",
        "4. Расширить OOS: 2024 holdout, 2026 H2 после появления данных, стресс-тест комиссий/проскальзывания.",
        "",
        "## Источники данных",
        "",
        "Метрики и графики построены из локальных артефактов в `Results/*` и сохранены в текущей папке отчёта.",
        "",
    ]
    md.write_text("\n".join(lines), encoding="utf-8")
    return md


def build_docx(rows: list[dict[str, Any]]) -> Path:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("python-docx is required for DOCX generation") from exc

    final_2025 = next(row for row in rows if row["period"] == "2025" and row["name"] == "round2_plus_terminal_atom")
    final_2026 = next(row for row in rows if row["period"] == "2026 H1" and row["name"] == "round2_plus_terminal_atom")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for style_name, size, color in [
        ("Title", 20, RGBColor(11, 37, 69)),
        ("Heading 1", 15, RGBColor(46, 116, 181)),
        ("Heading 2", 12, RGBColor(31, 77, 120)),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = color

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Panteon Flash: итоговый отчёт по готовности к реальным торгам")
    subtitle = doc.add_paragraph("Дата: 2026-05-23 | Ветка анализа: Panteon_Flash")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "Лучший устойчивый вариант для следующего этапа: "
        "proven_solo_portfolio_hard_solo_momentum_cap10_generated_selected_deny_round2_terminal_atom."
    )
    doc.add_paragraph(
        "Результат достаточен для ограниченного pilot-режима, но не для неограниченного запуска реальных торгов. "
        "Обязательны малые лимиты, shadow-мониторинг, daily attribution и автоматический kill-switch."
    )

    doc.add_heading("Лучший результат", level=1)
    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ["Период", "Flash PnL", "Best comp", "Alpha", "Beats best", "LCB", "Churn/day", "Exec signals"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for row in [final_2025, final_2026]:
        cells = table.add_row().cells
        values = [
            row["period"],
            _fmt(row["pnl_pct"]) + "%",
            _fmt(row["best_component_pnl_pct"]) + "%",
            _fmt(row["panteon_alpha_pct"]) + " п.п.",
            _fmt(row["beats_best_component"]),
            _fmt(row["panteon_lcb_dominance"]),
            _fmt(row["actor_switches_per_day"]),
            str(row["flash_executable_selected_signals"]),
        ]
        for cell, value in zip(cells, values):
            cell.text = value

    doc.add_paragraph(
        "Конфигурация: базовый cap10 portfolio, generated selected-deny из 2025 base+v1, "
        "и terminal guard для ensemble:Solo_LiveCrashHunter|ATOM/USDT|FUT_SHORT_FULL."
    )

    doc.add_heading("Визуализации", level=1)
    for filename, caption in [
        ("01_pnl_vs_best_component.png", "Flash против лучшего standalone-компонента."),
        ("02_alpha_by_iteration.png", "Alpha по итерациям и strict alpha threshold."),
        ("03_signals_and_churn.png", "Активность, исполнимые сигналы и churn."),
        ("04_gate_matrix.png", "Production-gates по ключевым итерациям."),
    ]:
        doc.add_picture(str(CHARTS / filename), width=Inches(6.3))
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Причины успехов", level=1)
    for text in [
        "Контроль конкретных actor|symbol|action ключей оказался полезнее, чем запрет акторов целиком.",
        "Round2 selected-deny снял selection bias: 2025 вырос с 3.76% до 21.35%.",
        "Terminal ATOM guard улучшил 2026 H1, если применять его поверх round2, а не вместо round2.",
        "Финальный вариант прошёл оба контрольных периода: 2025 остался выше лучшего компонента, а 2026 H1 преодолел strict alpha threshold.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("Причины провалов", level=1)
    for text in [
        "Broad shadow deny переобучался и ломал 2025: широкая фильтрация full-shadow статистики перерезала полезные ключи и снижала Flash PnL до 2.57%.",
        "Одиночный ATOM deny без round2-базы выглядел хорошо на 2026 H1, но проваливал 2025: без уже обученного round2 слоя Flash начинал выбирать плохие fallback-кандидаты.",
        "Расширение open-limit до open2 ухудшало OOS: 2026 H1 снижался с 10.14% до 9.54%, значит текущие лимиты реально защищают от плохой экспозиции.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("Рекомендация по реальным торгам", level=1)
    doc.add_paragraph(
        "Запускать можно только ограниченный pilot: сначала paper/live-shadow, затем малый real capital. "
        "Полноценный production запрещён до полного 2022-2026 walk-forward и стресс-теста комиссий."
    )
    for text in [
        "max_new_opens_per_bar оставить 1;",
        "kill-switch при rolling PanteonAdvantage < 0 или drawdown > 1.5-2.0%;",
        "ежедневно проверять standalone_vs_flash_selected и flash_attribution_summary;",
        "не добавлять новые deny-правила без 2025 + 2026 H1 + holdout валидации.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("Артефакты", level=1)
    for artifact in [
        OUT / "panteon_flash_iteration_metrics.csv",
        OUT / "panteon_flash_iteration_metrics.json",
        OUT / "Panteon_Flash_final_trading_readiness_report.md",
    ]:
        doc.add_paragraph(str(artifact), style="List Bullet")

    out = OUT / "Panteon_Flash_final_trading_readiness_report.docx"
    doc.save(out)
    return out


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--no-charts", action="store_true")
    parser.add_argument("--no-docx", action="store_true")
    args = parser.parse_args()
    rows = build_data()
    if not args.no_charts:
        build_charts(rows)
    md = build_markdown(rows)
    print(f"markdown={md}")
    if not args.no_docx:
        docx = build_docx(rows)
        print(f"docx={docx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
