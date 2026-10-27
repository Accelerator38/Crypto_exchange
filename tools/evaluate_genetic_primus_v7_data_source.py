from __future__ import annotations

import argparse
import json
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from exia.genetic_primus.contracts import sha256_file  # noqa: E402
from exia.genetic_primus.contracts_v7 import load_v7_contract  # noqa: E402
from exia.genetic_primus.data_source_v7 import validate_v7_dataset_manifest  # noqa: E402


DEFAULT_CONTRACT = ROOT / "configs/genetic_primus_v7_contract.json"
DEFAULT_MANIFEST = ROOT / "Retrodate/genetic_primus_v7_coinm_liquidations_v1_20230625_20241014/manifest.json"
DEFAULT_OUTPUT = ROOT / "Reports/Exia/Genetic_Primus/v7_data_source_gate_20260824"


def _workbook(path: Path, frames: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.book[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
    workbook = load_workbook(path, read_only=False)
    for name, frame in frames.items():
        sheet = workbook[name]
        if sheet.max_column != len(frame.columns) or sheet.max_column <= 1 or sheet.merged_cells.ranges:
            raise RuntimeError(f"invalid V7 source-gate workbook: {name}")
    workbook.close()


def evaluate(contract_path: Path, manifest_path: Path, output_dir: Path) -> dict[str, Any]:
    contract = load_v7_contract(contract_path, root=ROOT)
    manifest = validate_v7_dataset_manifest(contract, manifest_path, root=ROOT)
    output_dir.mkdir(parents=True, exist_ok=True)
    checks: list[dict[str, Any]] = []

    def add(check_id: str, actual: Any, required: Any, passed: bool) -> None:
        checks.append({"check_id": check_id, "actual": actual, "required": required, "passed": bool(passed)})

    source = contract.liquidation_source
    add("dataset_status", manifest["status"], "V7_LIQUIDATION_DATASET_BUILT", manifest["status"] == "V7_LIQUIDATION_DATASET_BUILT")
    add("point_in_time", manifest["point_in_time"], True, manifest["point_in_time"] is True)
    add("official_checksums", manifest["official_checksums_verified"], True, manifest["official_checksums_verified"] is True)
    add("raw_archive", manifest["raw_archive_retained"], True, manifest["raw_archive_retained"] is True)
    add(
        "missing_archives_not_zero_imputed",
        manifest["missing_daily_archives_zero_imputed"],
        False,
        manifest["missing_daily_archives_zero_imputed"] is False,
    )
    add("official_file_coverage", float(manifest["official_file_coverage"]), 1.0, float(manifest["official_file_coverage"]) == 1.0)
    cross_required = float(source["minimum_usable_cross_market_bar_coverage"])
    add(
        "usable_cross_market_bar_coverage",
        float(manifest["usable_cross_market_bar_coverage"]),
        cross_required,
        float(manifest["usable_cross_market_bar_coverage"]) >= cross_required,
    )
    coverage_rows = []
    day_required = float(source["minimum_archive_day_coverage_per_symbol"])
    for row in manifest["coverage"]:
        coverage = float(row["archive_day_coverage"])
        passed = coverage >= day_required
        add(f"{row['symbol']}:archive_day_coverage", coverage, day_required, passed)
        coverage_rows.append({**row, "required": day_required, "passed": passed})
    failed = [row["check_id"] for row in checks if not row["passed"]]
    passed = not failed
    status = "V7_DATA_INTEGRITY_ACCEPTED_FEASIBILITY_PENDING" if passed else "V7_TERMINAL_REJECTED_DATA_SOURCE"
    report = {
        "schema_version": "exia.genetic_primus_data_source_gate.v7",
        "generated_at": datetime.now(UTC).isoformat(),
        "status": status,
        "contract_sha256": contract.sha256,
        "dataset_manifest_sha256": sha256_file(manifest_path),
        "dataset_sha256": manifest["dataset_sha256"],
        "checks": checks,
        "failed_checks": failed,
        "integrity_and_coverage_passed": passed,
        "feasibility_allowed": passed,
        "feasibility_opened": False,
        "evolution_allowed": False,
        "validation_opened": False,
        "oos_opened": False,
        "paper_allowed": False,
        "live_allowed": False,
        "orders_enabled": False,
        "promotion_authority": False,
    }
    report_path = output_dir / "report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    _workbook(output_dir / "data_source_gate.xlsx", {"Checks": pd.DataFrame(checks), "Coverage": pd.DataFrame(coverage_rows)})
    lines = [
        "# Genetic_Primus V7 data-source gate",
        "",
        f"Status: `{status}`.",
        "",
        f"Official files: {manifest['downloaded_raw_files']}/{manifest['listed_raw_files']}.",
        f"Feature rows: {manifest['feature_rows']}.",
        f"Usable cross-market coverage: {manifest['usable_cross_market_bar_coverage']:.2%}.",
        "",
        "Event feasibility may be opened exactly once." if passed else "V7 is terminally blocked before future-return labels.",
    ]
    (output_dir / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return report


def assert_feasibility_authority(report_path: Path) -> dict[str, Any]:
    report = json.loads(report_path.read_text(encoding="utf-8"))
    if report.get("status") != "V7_DATA_INTEGRITY_ACCEPTED_FEASIBILITY_PENDING":
        raise PermissionError(f"V7 feasibility blocked by source status: {report.get('status')}")
    if report.get("feasibility_allowed") is not True or report.get("feasibility_opened") is not False:
        raise PermissionError("V7 feasibility authority is absent or consumed")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Evaluate Genetic_Primus V7 liquidation source")
    parser.add_argument("--contract", type=Path, default=DEFAULT_CONTRACT)
    parser.add_argument("--dataset-manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = evaluate(args.contract.resolve(), args.dataset_manifest.resolve(), args.output_dir.resolve())
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
