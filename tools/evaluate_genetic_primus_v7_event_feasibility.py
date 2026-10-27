from __future__ import annotations

import argparse
import json
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))
TOOLS = ROOT / "tools"
if str(TOOLS) not in sys.path:
    sys.path.insert(0, str(TOOLS))

from evaluate_genetic_primus_v7_data_source import assert_feasibility_authority  # noqa: E402
from exia.genetic_primus.contracts import sha256_file  # noqa: E402
from exia.genetic_primus.contracts_v7 import load_v7_contract  # noqa: E402
from exia.genetic_primus.data_source_v7 import validate_v7_dataset_manifest  # noqa: E402
from exia.genetic_primus.feasibility_v7 import build_event_ledger, score_events, slice_metrics  # noqa: E402
from exia.genetic_primus.features import build_feature_frame  # noqa: E402


DEFAULT_CONTRACT = ROOT / "configs/genetic_primus_v7_contract.json"
DEFAULT_DATASET_MANIFEST = ROOT / "Retrodate/genetic_primus_v7_coinm_liquidations_v1_20230625_20241014/manifest.json"
DEFAULT_SOURCE_GATE = ROOT / "Reports/Exia/Genetic_Primus/v7_data_source_gate_20260824/report.json"
DEFAULT_OUTPUT = ROOT / "Reports/Exia/Genetic_Primus/v7_event_feasibility_20260824"


def _load_execution_panel(contract) -> pd.DataFrame:
    source = contract.payload["execution_source"]
    manifest_path = ROOT / source["manifest"]
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    frames = []
    for item in manifest["files"]:
        path = ROOT / str(item["path"]).replace("\\", "/")
        if sha256_file(path) != item["sha256"]:
            raise ValueError("V7 execution panel artifact SHA mismatch")
        frames.append(pd.read_parquet(path))
    panel = pd.concat(frames, ignore_index=True)
    start = int(pd.Timestamp(contract.liquidation_source["start_date"], tz="UTC").timestamp() * 1000)
    end = int((pd.Timestamp(contract.liquidation_source["end_date"], tz="UTC") + pd.Timedelta(days=2)).timestamp() * 1000)
    return panel.loc[(panel["timestamp"] >= start) & (panel["timestamp"] < end)].copy()


def _workbook(path: Path, frames: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.book[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
    workbook = load_workbook(path, read_only=False)
    for name, frame in frames.items():
        sheet = workbook[name]
        if sheet.max_column != len(frame.columns) or sheet.max_column <= 1 or sheet.merged_cells.ranges:
            raise RuntimeError(f"invalid V7 feasibility workbook: {name}")
    workbook.close()


def _plot(folds: pd.DataFrame, path: Path) -> None:
    figure, axes = plt.subplots(1, 2, figsize=(14, 5), constrained_layout=True)
    for event_id, subset in folds.groupby("event_id", sort=True):
        axes[0].plot(subset["fold"], subset["mean_net_bps"], marker="o", label=event_id)
        axes[1].plot(subset["fold"], subset["lcb_net_bps"], marker="o", label=event_id)
    for axis, title in zip(axes, ("V7 stress-costed fold mean", "V7 stress-costed fold LCB"), strict=True):
        axis.axhline(0.0, color="black", linewidth=1)
        axis.set_title(title)
        axis.set_ylabel("bps per event")
        axis.grid(alpha=0.25)
        axis.tick_params(axis="x", rotation=25)
    axes[0].legend(fontsize=7)
    figure.savefig(path, dpi=160)
    plt.close(figure)


def evaluate(contract_path: Path, dataset_manifest_path: Path, source_gate_path: Path, output_dir: Path) -> dict[str, Any]:
    authority = assert_feasibility_authority(source_gate_path)
    contract = load_v7_contract(contract_path, root=ROOT)
    dataset_manifest = validate_v7_dataset_manifest(contract, dataset_manifest_path, root=ROOT)
    if authority["contract_sha256"] != contract.sha256 or authority["dataset_sha256"] != dataset_manifest["dataset_sha256"]:
        raise ValueError("V7 feasibility authority does not pin the current contract/dataset")
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = output_dir / "manifest.json"
    if manifest_path.is_file():
        return json.loads(manifest_path.read_text(encoding="utf-8"))
    panel = build_feature_frame(_load_execution_panel(contract))
    panel["market_return_1"] = panel.groupby("timestamp", sort=True)["return_1"].transform("mean")
    source = pd.read_parquet(ROOT / dataset_manifest["artifacts"]["feature_panel"]["path"])
    joined = panel.merge(source, on=["timestamp", "symbol"], how="inner", validate="one_to_one")
    ledger = build_event_ledger(joined, contract.feasibility, int(contract.payload["execution_source"]["holding_bars"]))
    summary, folds = score_events(ledger, folds=contract.folds, config=contract.feasibility)
    accepted_parts = []
    for fold in contract.folds:
        start = int(pd.Timestamp(fold["evaluation_start"]).timestamp() * 1000)
        end = int(pd.Timestamp(fold["evaluation_end"]).timestamp() * 1000)
        accepted_parts.append(
            ledger.loc[
                (ledger["timestamp"] >= start)
                & (ledger["timestamp"] < end)
                & (ledger["exit_timestamp"] < end)
            ]
        )
    accepted = pd.concat(accepted_parts, ignore_index=True)
    regimes = slice_metrics(accepted, contract.feasibility, "market_state")
    directions = slice_metrics(accepted, contract.feasibility, "direction")
    symbols = slice_metrics(accepted, contract.feasibility, "symbol")
    ledger.to_parquet(output_dir / "events.parquet", index=False)
    summary.to_parquet(output_dir / "summary.parquet", index=False)
    folds.to_parquet(output_dir / "folds.parquet", index=False)
    _workbook(
        output_dir / "event_feasibility.xlsx",
        {"Summary": summary, "Folds": folds, "Regimes": regimes, "Directions": directions, "Symbols": symbols, "Events": ledger},
    )
    _plot(folds, output_dir / "event_feasibility.png")
    passing = summary.loc[summary["passed"], "event_id"].astype(str).tolist()
    status = "V7_DATA_SOURCE_AND_EVENT_CAPACITY_ACCEPTED" if passing else "V7_TERMINAL_REJECTED_NO_EVENT_CAPACITY"
    manifest = {
        "schema_version": "exia.genetic_primus_event_feasibility.v7",
        "generated_at": datetime.now(UTC).isoformat(),
        "status": status,
        "terminal_reason": None if passing else "NO_PREREGISTERED_EVENT_HAS_POSITIVE_STRESS_COSTED_WALK_FORWARD_CAPACITY",
        "contract_sha256": contract.sha256,
        "dataset_manifest_sha256": sha256_file(dataset_manifest_path),
        "dataset_sha256": dataset_manifest["dataset_sha256"],
        "source_gate_sha256": sha256_file(source_gate_path),
        "passing_event_ids": passing,
        "event_summary": summary.to_dict(orient="records"),
        "labels_opened": True,
        "evolution_allowed": bool(passing),
        "validation_opened": False,
        "oos_opened": False,
        "paper_allowed": False,
        "live_allowed": False,
        "orders_enabled": False,
        "promotion_authority": False,
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    lines = [
        "# Genetic_Primus V7 liquidation event feasibility",
        "",
        f"Status: `{status}`.",
        "",
        "| Event | Events | Mean | Median | LCB | Positive folds | Worst fold mean | Median fold LCB | Passed |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---|",
    ]
    for row in summary.itertuples(index=False):
        lines.append(
            f"| {row.event_id} | {row.events} | {row.mean_net_bps:.2f} | {row.median_net_bps:.2f} | "
            f"{row.lcb_net_bps:.2f} | {row.positive_folds}/5 | {row.worst_fold_mean_bps:.2f} | "
            f"{row.median_fold_lcb_bps:.2f} | {row.passed} |"
        )
    lines.extend(["", "Evolution remains blocked." if not passing else "Passing liquidation events may open fixed V7 model screening."])
    (output_dir / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    freeze = {
        "schema_version": "exia.genetic_primus_terminal_freeze.v7",
        "generated_at": manifest["generated_at"],
        "status": status,
        "terminal_reason": manifest["terminal_reason"],
        "contract_sha256": contract.sha256,
        "dataset_sha256": dataset_manifest["dataset_sha256"],
        "event_manifest_sha256": sha256_file(manifest_path),
        "evolution_allowed": bool(passing),
        "validation_allowed": False,
        "oos_allowed": False,
        "paper_allowed": False,
        "live_allowed": False,
        "orders_enabled": False,
        "promotion_authority": False,
    }
    (output_dir / "freeze.json").write_text(json.dumps(freeze, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description="Run preregistered Genetic_Primus V7 liquidation feasibility")
    parser.add_argument("--contract", type=Path, default=DEFAULT_CONTRACT)
    parser.add_argument("--dataset-manifest", type=Path, default=DEFAULT_DATASET_MANIFEST)
    parser.add_argument("--source-gate", type=Path, default=DEFAULT_SOURCE_GATE)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = evaluate(args.contract.resolve(), args.dataset_manifest.resolve(), args.source_gate.resolve(), args.output_dir.resolve())
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
