from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


WORKSPACE = Path(r"C:\Work\Crypto_exchange")
BASELINE = WORKSPACE / (
    r"Results\neiro_genetics\RetrodateMarket_AB\C_post_registry_strict"
    r"\RETRODATE_MARKET\2026-05-19_14-32-38_retrodate_market_v2"
)
FINAL = WORKSPACE / (
    r"Results\neiro_genetics\RetrodateMarket_AB\Q_final_scoped_probation_2022_2026"
    r"\RETRODATE_MARKET\2026-05-19_23-30-20_retrodate_market_v2"
)
DIAGNOSTICS = {
    "strict_old": WORKSPACE / (
        r"Results\neiro_genetics\RetrodateMarket_AB\K_diag_antonius_gate_5100"
        r"\RETRODATE_MARKET\2026-05-19_22-39-42_retrodate_market_v2"
    ),
    "wide_agent_fallback": WORKSPACE / (
        r"Results\neiro_genetics\RetrodateMarket_AB\L_diag_agent_registry_fallback_5100"
        r"\RETRODATE_MARKET\2026-05-19_22-50-39_retrodate_market_v2"
    ),
    "no_maturity_bar": WORKSPACE / (
        r"Results\neiro_genetics\RetrodateMarket_AB\M_diag_agent_fallback_no_maturity_bar_5100"
        r"\RETRODATE_MARKET\2026-05-19_23-00-00_retrodate_market_v2"
    ),
    "recent_deny": WORKSPACE / (
        r"Results\neiro_genetics\RetrodateMarket_AB\N_diag_filled_recent_deny_5100"
        r"\RETRODATE_MARKET\2026-05-19_23-08-48_retrodate_market_v2"
    ),
    "probation_min1": WORKSPACE / (
        r"Results\neiro_genetics\RetrodateMarket_AB\O_diag_probation_min1_5100"
        r"\RETRODATE_MARKET\2026-05-19_23-15-49_retrodate_market_v2"
    ),
    "scoped_probation": WORKSPACE / (
        r"Results\neiro_genetics\RetrodateMarket_AB\P_diag_scoped_probation_5100"
        r"\RETRODATE_MARKET\2026-05-19_23-23-05_retrodate_market_v2"
    ),
}
OUT_DIR = WORKSPACE / (
    r"Results\neiro_genetics\RetrodateMarket_AB"
    r"\analysis_antonius_static_rotator_2026_05_20"
)


def font(size: int, *, bold: bool = False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    path = Path(r"C:\Windows\Fonts") / name
    try:
        return ImageFont.truetype(str(path), size=size)
    except Exception:
        return ImageFont.load_default()


@dataclass(frozen=True)
class RunMetrics:
    label: str
    path: Path
    pnl_pct: float
    pnl_usd: float
    max_dd_pct: float
    real_closed_trades: int
    no_trade_pct: float
    raw_zero_pct: float
    filled_zero_pct: float
    best_single_label: str
    best_single_usd: float
    best_soft_label: str
    best_soft_usd: float
    perfect_usd: float
    perfect_pct: float
    perfect_dd_pct: float
    mismatch_counts: dict[str, int]


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def extract_float(pattern: str, text: str, default: float = 0.0) -> float:
    match = re.search(pattern, text)
    if not match:
        return default
    return float(match.group(1).replace(",", ""))


def extract_int(pattern: str, text: str, default: int = 0) -> int:
    match = re.search(pattern, text)
    if not match:
        return default
    return int(float(match.group(1).replace(",", "")))


def load_run(label: str, path: Path) -> RunMetrics:
    report = read_text(path / "analysis_report.md")
    soft = json.loads((path / "soft_allocator_report.json").read_text(encoding="utf-8"))
    perfect = json.loads((path / "perfect_panteon_report.json").read_text(encoding="utf-8"))
    allocation = json.loads((path / "allocation_diagnostics.json").read_text(encoding="utf-8"))
    oracle = json.loads((path / "oracle_mismatch_report.json").read_text(encoding="utf-8"))
    return RunMetrics(
        label=label,
        path=path,
        pnl_pct=extract_float(r"Panteon owned PnL: ([\-0-9.]+)%", report),
        pnl_usd=extract_float(r"Panteon realized PnL USD: \$([\-0-9.]+)", report),
        max_dd_pct=extract_float(r"Panteon max drawdown: ([0-9.]+)%", report),
        real_closed_trades=extract_int(r"Real closed trades: ([0-9.]+)", report),
        no_trade_pct=float(allocation.get("no_trade_share_pct", 0.0)),
        raw_zero_pct=float(allocation.get("raw_zero_share_pct", 0.0)),
        filled_zero_pct=float(allocation.get("filled_zero_share_pct", 0.0)),
        best_single_label=str(soft.get("best_single_label", "")),
        best_single_usd=float(soft.get("best_single_pnl_usd", 0.0)),
        best_soft_label=str(soft.get("best_policy_name", "")),
        best_soft_usd=float(soft.get("best_policy_pnl_usd", 0.0)),
        perfect_usd=float(perfect.get("pnl_usd", 0.0)),
        perfect_pct=float(perfect.get("pnl_pct", 0.0)),
        perfect_dd_pct=float(perfect.get("max_drawdown_pct", 0.0)),
        mismatch_counts=dict(oracle.get("summary", {}).get("reason_counts", {})),
    )


def real_contributors(path: Path) -> list[tuple[str, float, int, float, float]]:
    data = json.loads((path / "leaderboard_players.json").read_text(encoding="utf-8"))
    rows = []
    for label, payload in data.get("players", {}).items():
        pnl = float(payload.get("realized_pnl_usd", 0.0) or 0.0)
        trades = int(payload.get("real_trades", 0) or 0)
        if trades <= 0 and abs(pnl) < 1e-12:
            continue
        rows.append((
            label.replace("V_", ""),
            pnl,
            trades,
            float(payload.get("real_win_rate", 0.0) or 0.0),
            float(payload.get("pnl_pct", 0.0) or 0.0),
        ))
    rows.sort(key=lambda item: item[1], reverse=True)
    return rows


def top_virtual(path: Path, filename: str, limit: int = 8) -> list[tuple[str, float, int, float, float]]:
    data = json.loads((path / filename).read_text(encoding="utf-8"))
    key = "agents" if "agents" in data else "players"
    rows = []
    for label, payload in data.get(key, {}).items():
        rows.append((
            label.replace("V_", ""),
            float(payload.get("pnl_pct", 0.0) or 0.0),
            int(payload.get("closed_trades", 0) or 0),
            float(payload.get("win_rate", 0.0) or 0.0),
            float(payload.get("max_drawdown_pct", 0.0) or 0.0),
        ))
    rows.sort(key=lambda item: item[1], reverse=True)
    return rows[:limit]


def draw_bar_chart(
    path: Path,
    title: str,
    rows: list[tuple[str, float]],
    *,
    suffix: str = "",
    width: int = 1300,
    height: int = 620,
) -> None:
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    body_font = font(18)
    title_font = font(22, bold=True)
    left, top, right, bottom = 260, 90, width - 90, height - 110
    draw.text((40, 28), title, fill=(20, 34, 52), font=title_font)
    if not rows:
        img.save(path)
        return
    values = [value for _label, value in rows]
    max_abs = max(abs(v) for v in values) or 1.0
    zero_x = left + (right - left) / 2
    if all(v >= 0 for v in values):
        zero_x = left
    bar_h = max(18, int((bottom - top) / max(1, len(rows)) * 0.55))
    step = (bottom - top) / max(1, len(rows))
    draw.line((left, bottom, right, bottom), fill=(180, 188, 196), width=1)
    draw.line((zero_x, top - 8, zero_x, bottom + 8), fill=(140, 148, 156), width=1)
    for idx, (label, value) in enumerate(rows):
        y = top + idx * step + step * 0.25
        scale = (right - left) / (max_abs if all(v >= 0 for v in values) else max_abs * 2)
        x2 = zero_x + value * scale
        color = (31, 122, 87) if value >= 0 else (176, 58, 46)
        draw.rectangle((min(zero_x, x2), y, max(zero_x, x2), y + bar_h), fill=color)
        draw.text((30, y + 2), label[:28], fill=(35, 45, 55), font=body_font)
        value_text = f"{value:.2f}{suffix}"
        text_x = max(zero_x, x2) + 10
        bbox = draw.textbbox((0, 0), value_text, font=body_font)
        if text_x + (bbox[2] - bbox[0]) > width - 25:
            text_x = min(zero_x, x2) + 10
            fill = "white" if abs(x2 - zero_x) > 160 else (35, 45, 55)
        else:
            fill = (35, 45, 55)
        draw.text((text_x, y + 2), value_text, fill=fill, font=body_font)
    img.save(path)


def draw_metric_table_chart(path: Path, baseline: RunMetrics, final: RunMetrics) -> None:
    rows = [
        ("PnL USD", baseline.pnl_usd, final.pnl_usd),
        ("PnL %", baseline.pnl_pct, final.pnl_pct),
        ("Max DD %", baseline.max_dd_pct, final.max_dd_pct),
        ("NoTrade %", baseline.no_trade_pct, final.no_trade_pct),
        ("Raw-zero %", baseline.raw_zero_pct, final.raw_zero_pct),
        ("Trades", baseline.real_closed_trades, final.real_closed_trades),
    ]
    img = Image.new("RGB", (1100, 520), "white")
    draw = ImageDraw.Draw(img)
    body_font = font(17)
    title_font = font(22, bold=True)
    draw.text((40, 30), "Baseline vs final run", fill=(20, 34, 52), font=title_font)
    x = [40, 350, 565, 780, 960]
    y0 = 90
    headers = ["Metric", "Baseline", "Final", "Delta", "Direction"]
    for i, header in enumerate(headers):
        draw.text((x[i], y0), header, fill=(20, 34, 52), font=body_font)
    draw.line((35, y0 + 22, 1060, y0 + 22), fill=(160, 168, 176), width=1)
    for idx, (name, base, fin) in enumerate(rows, start=1):
        y = y0 + idx * 55
        delta = fin - base
        if name in {"Max DD %", "NoTrade %", "Raw-zero %"}:
            direction = "better" if delta < 0 else "worse"
        else:
            direction = "better" if delta > 0 else "worse"
        color = (31, 122, 87) if direction == "better" else (176, 58, 46)
        for i, value in enumerate([name, f"{base:.2f}", f"{fin:.2f}", f"{delta:+.2f}", direction]):
            draw.text((x[i], y), value, fill=color if i == 4 else (35, 45, 55), font=body_font)
    img.save(path)


def table_md(headers: list[str], rows: Iterable[Iterable[object]]) -> str:
    out = ["| " + " | ".join(headers) + " |"]
    out.append("| " + " | ".join(["---"] * len(headers)) + " |")
    for row in rows:
        out.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(out)


def build_markdown(
    out_dir: Path,
    baseline: RunMetrics,
    final: RunMetrics,
    diag: dict[str, RunMetrics],
    contributors: list[tuple[str, float, int, float, float]],
) -> str:
    positives = [row for row in contributors if row[1] > 0]
    profitable_share = len(positives) / len(contributors) * 100.0 if contributors else 0.0
    trade_profitable_share = (
        sum(row[2] for row in positives) / max(1, sum(row[2] for row in contributors)) * 100.0
    )
    lines = [
        "# Panteon v3 allocator: Antonius/static rotator iteration",
        "",
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        "",
        "## Executive summary",
        "",
        (
            f"Final 2022-2026 run improved realized PnL from ${baseline.pnl_usd:.2f} "
            f"to ${final.pnl_usd:.2f}, but the model is still negative and remains far "
            "from both soft allocator and Perfect Panteon."
        ),
        (
            f"Perfect monthly oracle is ${final.perfect_usd:.2f} ({final.perfect_pct:.2f}%), "
            f"while real Panteon is ${final.pnl_usd:.2f}. The remaining gap is "
            f"${final.perfect_usd - final.pnl_usd:.2f}."
        ),
        (
            f"Only {len(positives)}/{len(contributors)} real leaders were profitable "
            f"({profitable_share:.1f}% labels; {trade_profitable_share:.1f}% of real trades)."
        ),
        "",
        "## What changed",
        "",
        "- Fixed Antonius mapping and protected composite/regime-switch labels from hopeless quarantine.",
        "- Added `Optimal_StaticRotator`, a static agent-set player with regime-specific rotation order.",
        "- Added current-bar agent signal registry and fallback to executable `Solo_<agent>` signals.",
        "- Changed fallback registry to filled-only shadow signals and recent entry-causal ranking.",
        "- Moved retrotest experimental maturity away from fixed `bar>=5000`; shadow quality gates remain.",
        "- Added benchmark hard policy denylist for old weak profiles, fixed sets, Perfect test players, and weak solo variants.",
        "- Used scoped probation loss-kill after the first losing trade for fragile experimental labels.",
        "",
        "## Baseline vs final",
        "",
        table_md(
            ["Metric", "Baseline strict", "Final scoped", "Delta"],
            [
                ["PnL USD", f"{baseline.pnl_usd:.2f}", f"{final.pnl_usd:.2f}", f"{final.pnl_usd - baseline.pnl_usd:+.2f}"],
                ["PnL %", f"{baseline.pnl_pct:.2f}", f"{final.pnl_pct:.2f}", f"{final.pnl_pct - baseline.pnl_pct:+.2f} pp"],
                ["Max DD %", f"{baseline.max_dd_pct:.2f}", f"{final.max_dd_pct:.2f}", f"{final.max_dd_pct - baseline.max_dd_pct:+.2f} pp"],
                ["NoTrade %", f"{baseline.no_trade_pct:.2f}", f"{final.no_trade_pct:.2f}", f"{final.no_trade_pct - baseline.no_trade_pct:+.2f} pp"],
                ["Raw-zero %", f"{baseline.raw_zero_pct:.2f}", f"{final.raw_zero_pct:.2f}", f"{final.raw_zero_pct - baseline.raw_zero_pct:+.2f} pp"],
                ["Real trades", baseline.real_closed_trades, final.real_closed_trades, f"{final.real_closed_trades - baseline.real_closed_trades:+d}"],
            ],
        ),
        "",
        "![Baseline vs final](metric_comparison.png)",
        "",
        "## Short diagnostic A/B",
        "",
        table_md(
            ["Run", "PnL %", "PnL USD", "DD %", "Trades", "NoTrade %"],
            [
                [key, f"{run.pnl_pct:.2f}", f"{run.pnl_usd:.2f}", f"{run.max_dd_pct:.2f}", run.real_closed_trades, f"{run.no_trade_pct:.2f}"]
                for key, run in diag.items()
            ],
        ),
        "",
        "## Potential gap",
        "",
        table_md(
            ["Reference", "USD", "Pct/DD", "Comment"],
            [
                ["Real Panteon", f"{final.pnl_usd:.2f}", f"{final.pnl_pct:.2f}% / DD {final.max_dd_pct:.2f}%", "actual executable result"],
                [f"Best single shadow ({final.best_single_label})", f"{final.best_single_usd:.2f}", "-", "not executable as-is without position handoff"],
                [f"Best soft policy ({final.best_soft_label})", f"{final.best_soft_usd:.2f}", "-", "offline allocator potential"],
                ["Perfect monthly oracle", f"{final.perfect_usd:.2f}", f"{final.perfect_pct:.2f}% / DD {final.perfect_dd_pct:.2f}%", "upper bound target"],
            ],
        ),
        "",
        "![Potential gap](potential_gap.png)",
        "",
        "## Real leader contribution",
        "",
        table_md(
            ["Leader", "Realized USD", "Trades", "Real win %", "Shadow PnL %"],
            [
                [label, f"{pnl:.2f}", trades, f"{win:.1f}", f"{shadow:.2f}"]
                for label, pnl, trades, win, shadow in contributors
            ],
        ),
        "",
        "![Real contribution](real_contribution.png)",
        "",
        "## Main diagnosis",
        "",
        "- The code package reduced loss and removed several damaging real leaders, but it mostly achieved this by suppressing trading.",
        f"- NoTrade remains extreme: {final.no_trade_pct:.2f}% of bars. The allocator is not yet converting shadow edge into executable entries.",
        "- Best virtual agents/players are profitable, but their profit comes from shadow position lifecycle. Real fallback often sees the signal after the profitable shadow entry has already happened.",
        "- Antonius_conservative is strong in shadow (+62.23% player PnL), but direct real execution lost on the first trade and was killed by probation. This is a causality/handoff issue, not simply a scoring issue.",
        "- Optimal_StaticRotator exists and is visible, but its shadow PnL was negative on the full run. Its current static order is not yet the optimal production static player.",
        "- Perfect Panteon rose to +478.27%, so the ceiling is high. The gap is mainly allocator/execution causality, not absence of profitable components.",
        "",
        "## Recommendations",
        "",
        "1. Do not return this Panteon version to live trading yet. It is safer than previous broad fallback variants, but still negative and too inactive.",
        "2. Next technical step: implement shadow-position handoff/replay for selected leaders, so real execution can join a still-valid shadow position instead of waiting for a new raw signal.",
        "3. Build a causal entry dataset: each real decision must be trained/evaluated only on signals available before the entry, not on aggregate shadow PnL.",
        "4. Split agents into two classes: executable-entry agents and shadow-state agents. FundingArb/LiveCrashHunter/ResearchValidator show shadow edge, but need handoff logic before real use.",
        "5. Rebuild Optimal_StaticRotator from profitable executable leaders only: VolBreakoutHunter, LiveVolCompress, CrashPanicShortAgent are currently the only positive real contributors in this run.",
        "6. Keep Perfect players out of real selection; they are oracle/test targets, not production leaders.",
        "",
        "## Artifacts",
        "",
        f"- Final run: `{final.path}`",
        f"- Baseline run: `{baseline.path}`",
    ]
    return "\n".join(lines) + "\n"


def add_table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = str(header)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)


def build_docx(
    path: Path,
    baseline: RunMetrics,
    final: RunMetrics,
    diag: dict[str, RunMetrics],
    contributors: list[tuple[str, float, int, float, float]],
    charts: dict[str, Path],
) -> None:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    title = document.add_heading("Panteon v3 allocator: Antonius/static rotator iteration", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("5-year retrotest report, 2022-2026")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        f"Final run improved PnL from ${baseline.pnl_usd:.2f} to ${final.pnl_usd:.2f}, "
        "but Panteon remains negative and far from soft/perfect potential."
    )
    document.add_heading("Baseline vs final", level=1)
    add_table(document, ["Metric", "Baseline", "Final", "Delta"], [
        ["PnL USD", f"{baseline.pnl_usd:.2f}", f"{final.pnl_usd:.2f}", f"{final.pnl_usd - baseline.pnl_usd:+.2f}"],
        ["PnL %", f"{baseline.pnl_pct:.2f}", f"{final.pnl_pct:.2f}", f"{final.pnl_pct - baseline.pnl_pct:+.2f} pp"],
        ["Max DD %", f"{baseline.max_dd_pct:.2f}", f"{final.max_dd_pct:.2f}", f"{final.max_dd_pct - baseline.max_dd_pct:+.2f} pp"],
        ["NoTrade %", f"{baseline.no_trade_pct:.2f}", f"{final.no_trade_pct:.2f}", f"{final.no_trade_pct - baseline.no_trade_pct:+.2f} pp"],
        ["Raw-zero %", f"{baseline.raw_zero_pct:.2f}", f"{final.raw_zero_pct:.2f}", f"{final.raw_zero_pct - baseline.raw_zero_pct:+.2f} pp"],
        ["Real trades", baseline.real_closed_trades, final.real_closed_trades, final.real_closed_trades - baseline.real_closed_trades],
    ])
    document.add_picture(str(charts["metric"]), width=Inches(6.6))
    document.add_heading("Potential gap", level=1)
    add_table(document, ["Reference", "USD", "Notes"], [
        ["Real Panteon", f"{final.pnl_usd:.2f}", "actual executable result"],
        [f"Best single shadow ({final.best_single_label})", f"{final.best_single_usd:.2f}", "shadow potential"],
        [f"Best soft policy ({final.best_soft_label})", f"{final.best_soft_usd:.2f}", "offline allocator"],
        ["Perfect monthly oracle", f"{final.perfect_usd:.2f}", f"{final.perfect_pct:.2f}%"],
    ])
    document.add_picture(str(charts["gap"]), width=Inches(6.6))
    document.add_heading("Real leader contribution", level=1)
    add_table(document, ["Leader", "USD", "Trades", "Win %", "Shadow PnL %"], [
        [label, f"{pnl:.2f}", trades, f"{win:.1f}", f"{shadow:.2f}"]
        for label, pnl, trades, win, shadow in contributors
    ])
    document.add_picture(str(charts["contrib"]), width=Inches(6.6))
    document.add_heading("Short diagnostic A/B", level=1)
    add_table(document, ["Run", "PnL %", "USD", "DD %", "Trades", "NoTrade %"], [
        [key, f"{run.pnl_pct:.2f}", f"{run.pnl_usd:.2f}", f"{run.max_dd_pct:.2f}", run.real_closed_trades, f"{run.no_trade_pct:.2f}"]
        for key, run in diag.items()
    ])
    document.add_heading("Diagnosis", level=1)
    for item in [
        "The package reduced loss but did not create a profitable Panteon.",
        f"NoTrade is still {final.no_trade_pct:.2f}%: the allocator is not converting shadow edge into executable real entries.",
        "Best shadow players are profitable mostly because they own full shadow position lifecycle; real fallback enters too late.",
        "Antonius_conservative is strong in shadow, but lost in real on its first trade and was killed by probation.",
        "Optimal_StaticRotator is present but not yet a profitable full-period static player.",
    ]:
        document.add_paragraph(item, style=None)
    document.add_heading("Recommendations", level=1)
    for item in [
        "Do not return this Panteon version to live trading yet.",
        "Implement shadow-position handoff/replay as the next technical step.",
        "Build entry-causal training/evaluation data and stop relying on aggregate shadow PnL for real entries.",
        "Restrict production candidates to real-positive executable leaders until handoff exists.",
        "Keep Perfect players as oracle targets only, never as real leaders.",
    ]:
        document.add_paragraph(item, style=None)
    document.save(path)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    baseline = load_run("baseline_strict", BASELINE)
    final = load_run("final_scoped", FINAL)
    diag = {key: load_run(key, path) for key, path in DIAGNOSTICS.items()}
    contributors = real_contributors(FINAL)

    charts = {
        "metric": OUT_DIR / "metric_comparison.png",
        "gap": OUT_DIR / "potential_gap.png",
        "contrib": OUT_DIR / "real_contribution.png",
    }
    draw_metric_table_chart(charts["metric"], baseline, final)
    draw_bar_chart(charts["gap"], "Real vs shadow/oracle potential", [
        ("Real Panteon", final.pnl_usd),
        (f"Best single: {final.best_single_label}", final.best_single_usd),
        (f"Best soft: {final.best_soft_label}", final.best_soft_usd),
        ("Perfect monthly oracle", final.perfect_usd),
    ], suffix=" USD")
    draw_bar_chart(
        charts["contrib"],
        "Real leader contribution, final run",
        [(label, pnl) for label, pnl, *_ in contributors],
        suffix=" USD",
        height=760,
    )

    markdown = build_markdown(OUT_DIR, baseline, final, diag, contributors)
    md_path = OUT_DIR / "Panteon_v3_Antonius_StaticRotator_Report.md"
    md_path.write_text(markdown, encoding="utf-8")
    docx_path = OUT_DIR / "Panteon_v3_Antonius_StaticRotator_Report.docx"
    build_docx(docx_path, baseline, final, diag, contributors, charts)
    print(md_path)
    print(docx_path)


if __name__ == "__main__":
    main()
