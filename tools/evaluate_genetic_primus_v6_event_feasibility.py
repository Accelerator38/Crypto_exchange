from __future__ import annotations

import argparse
import json
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))
TOOLS = ROOT / "tools"
if str(TOOLS) not in sys.path:
    sys.path.insert(0, str(TOOLS))

from evaluate_genetic_primus_v6_data_source import assert_feasibility_authority  # noqa: E402
from exia.genetic_primus.contracts import sha256_file  # noqa: E402
from exia.genetic_primus.contracts_v6 import load_v6_contract  # noqa: E402
from exia.genetic_primus.data_source_v6 import validate_v6_dataset_manifest  # noqa: E402
from exia.genetic_primus.feasibility_v6 import build_event_ledger, score_events, slice_metrics  # noqa: E402
from exia.genetic_primus.features import build_feature_frame  # noqa: E402


DEFAULT_CONTRACT = ROOT / "configs/genetic_primus_v6_contract.json"
DEFAULT_DATASET_MANIFEST = ROOT / "Retrodate/genetic_primus_v6_oi_kline_flow_v1_20220101_20241231/manifest.json"
DEFAULT_SOURCE_GATE = ROOT / "Reports/Exia/Genetic_Primus/v6_data_source_gate_20260824/report.json"
DEFAULT_OUTPUT = ROOT / "Reports/Exia/Genetic_Primus/v6_event_feasibility_20260824"


def _load_execution_panel(contract) -> pd.DataFrame:
    manifest = json.loads((ROOT / contract.payload["execution_source"]["manifest"]).read_text(encoding="utf-8"))
    frames = []
    for item in manifest["files"]:
        path = ROOT / str(item["path"]).replace("\\", "/")
        if sha256_file(path) != item["sha256"]:
            raise ValueError("V6 execution panel artifact SHA mismatch")
        frames.append(pd.read_parquet(path))
    panel = pd.concat(frames, ignore_index=True)
    end = int(pd.Timestamp("2025-01-01T00:00:00Z").timestamp() * 1000)
    return panel.loc[panel["timestamp"] < end].sort_values(["timestamp", "symbol"], kind="stable").reset_index(drop=True)


def _workbook(path: Path, frames: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.book[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
    workbook = load_workbook(path, read_only=False)
    for name, frame in frames.items():
        sheet = workbook[name]
        if sheet.max_column != len(frame.columns) or sheet.max_column <= 1 or sheet.merged_cells.ranges:
            raise RuntimeError(f"invalid V6 feasibility workbook: {name}")
    workbook.close()


def _plot(folds: pd.DataFrame, path: Path) -> None:
    figure, axes = plt.subplots(1, 2, figsize=(13, 5), constrained_layout=True)
    for event_id, subset in folds.groupby("event_id", sort=True):
        axes[0].plot(subset["fold"], subset["mean_net_bps"], marker="o", label=event_id)
        axes[1].plot(subset["fold"], subset["lcb_net_bps"], marker="o", label=event_id)
    for axis, title in zip(axes, ("Stress-costed fold mean", "Stress-costed fold LCB"), strict=True):
        axis.axhline(0.0, color="black", linewidth=1)
        axis.set_title(title)
        axis.set_ylabel("bps per event")
        axis.grid(alpha=0.25)
        axis.tick_params(axis="x", rotation=25)
    axes[0].legend(fontsize=8)
    figure.savefig(path, dpi=160)
    plt.close(figure)


def evaluate(contract_path: Path, dataset_manifest_path: Path, source_gate_path: Path, output_dir: Path) -> dict[str, Any]:
    authority = assert_feasibility_authority(source_gate_path)
    contract = load_v6_contract(contract_path, root=ROOT)
    dataset_manifest = validate_v6_dataset_manifest(contract, dataset_manifest_path, root=ROOT)
    if authority["contract_sha256"] != contract.sha256 or authority["dataset_sha256"] != dataset_manifest["dataset_sha256"]:
        raise ValueError("V6 feasibility authority does not pin the current contract/dataset")
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = output_dir / "manifest.json"
    if manifest_path.is_file():
        return json.loads(manifest_path.read_text(encoding="utf-8"))
    panel = _load_execution_panel(contract)
    features = build_feature_frame(panel)
    derivatives = pd.read_parquet(ROOT / dataset_manifest["artifacts"]["combined_dataset"]["path"])
    joined = features.merge(
        derivatives[
            [
                "timestamp",
                "symbol",
                "oi_change_4h",
                "oi_change_12h",
                "oi_change_24h",
                "oi_value_change_4h",
                "account_long_short_log",
                "kline_taker_imbalance_4h",
            ]
        ],
        on=["timestamp", "symbol"],
        how="inner",
        validate="one_to_one",
    )
    ledger = build_event_ledger(joined, contract.feasibility)
    summary, folds = score_events(ledger, folds=contract.folds, config=contract.feasibility)
    evaluation_start = int(pd.Timestamp(contract.folds[0]["evaluation_start"]).timestamp() * 1000)
    evaluation_end = int(pd.Timestamp(contract.folds[-1]["evaluation_end"]).timestamp() * 1000)
    evaluation_ledger = ledger.loc[(ledger["timestamp"] >= evaluation_start) & (ledger["timestamp"] < evaluation_end)].copy()
    regimes = slice_metrics(evaluation_ledger, contract.feasibility, "market_state")
    directions = slice_metrics(evaluation_ledger, contract.feasibility, "direction")
    ledger.to_parquet(output_dir / "events.parquet", index=False)
    summary.to_parquet(output_dir / "summary.parquet", index=False)
    folds.to_parquet(output_dir / "folds.parquet", index=False)
    _workbook(
        output_dir / "event_feasibility.xlsx",
        {"Summary": summary, "Folds": folds, "Regimes": regimes, "Directions": directions, "Events": ledger},
    )
    _plot(folds, output_dir / "event_feasibility.png")
    passing = summary.loc[summary["passed"], "event_id"].astype(str).tolist()
    status = "V6_DATA_SOURCE_AND_EVENT_CAPACITY_ACCEPTED" if passing else "V6_TERMINAL_REJECTED_NO_EVENT_CAPACITY"
    manifest = {
        "schema_version": "exia.genetic_primus_event_feasibility.v6",
        "generated_at": datetime.now(UTC).isoformat(),
        "status": status,
        "terminal_reason": None if passing else "NO_PREREGISTERED_EVENT_HAS_POSITIVE_STRESS_COSTED_WALK_FORWARD_CAPACITY",
        "contract_sha256": contract.sha256,
        "dataset_manifest_sha256": sha256_file(dataset_manifest_path),
        "dataset_sha256": dataset_manifest["dataset_sha256"],
        "source_gate_sha256": sha256_file(source_gate_path),
        "passing_event_ids": passing,
        "event_summary": summary.to_dict(orient="records"),
        "labels_opened": True,
        "evolution_allowed": bool(passing),
        "validation_opened": False,
        "oos_opened": False,
        "paper_allowed": False,
        "live_allowed": False,
        "orders_enabled": False,
        "promotion_authority": False,
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    lines = ["# Genetic_Primus V6 event feasibility", "", f"Status: `{status}`.", "", "| Event | Events | Mean | Median | LCB | Positive folds | Worst fold mean | Median fold LCB | Passed |", "|---|---:|---:|---:|---:|---:|---:|---:|---|"]
    for row in summary.itertuples(index=False):
        lines.append(f"| {row.event_id} | {row.events} | {row.mean_net_bps:.2f} | {row.median_net_bps:.2f} | {row.lcb_net_bps:.2f} | {row.positive_folds}/4 | {row.worst_fold_mean_bps:.2f} | {row.median_fold_lcb_bps:.2f} | {row.passed} |")
    lines.extend(["", "Evolution remains blocked." if not passing else "The passing event set may open the fixed V6 evolution stage."])
    (output_dir / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    terminal = {
        "schema_version": "exia.genetic_primus_terminal_freeze.v6",
        "generated_at": manifest["generated_at"],
        "status": status,
        "terminal_reason": manifest["terminal_reason"],
        "contract_sha256": contract.sha256,
        "dataset_sha256": dataset_manifest["dataset_sha256"],
        "event_manifest_sha256": sha256_file(manifest_path),
        "evolution_allowed": bool(passing),
        "validation_allowed": False,
        "oos_allowed": False,
        "paper_allowed": False,
        "live_allowed": False,
        "orders_enabled": False,
        "promotion_authority": False,
    }
    (output_dir / "freeze.json").write_text(json.dumps(terminal, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description="Run preregistered Genetic_Primus V6 event feasibility")
    parser.add_argument("--contract", type=Path, default=DEFAULT_CONTRACT)
    parser.add_argument("--dataset-manifest", type=Path, default=DEFAULT_DATASET_MANIFEST)
    parser.add_argument("--source-gate", type=Path, default=DEFAULT_SOURCE_GATE)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    print(json.dumps(evaluate(args.contract.resolve(), args.dataset_manifest.resolve(), args.source_gate.resolve(), args.output_dir.resolve()), ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
