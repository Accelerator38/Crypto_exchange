from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Mapping

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .contracts import MARKET_STATES
from .evolution import Champion, EvolutionResult


def _write_frame_sheet(writer: pd.ExcelWriter, name: str, frame: pd.DataFrame) -> None:
    frame.to_excel(writer, sheet_name=name, index=False)
    sheet = writer.book[name]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        values = [str(cell.value) if cell.value is not None else "" for cell in column[:200]]
        width = min(42, max(10, max(len(value) for value in values) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width


def write_workbook(
    path: Path,
    *,
    generations: pd.DataFrame,
    periods: pd.DataFrame,
    slices: pd.DataFrame,
    trades: pd.DataFrame,
) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _write_frame_sheet(writer, "Generations", generations)
        _write_frame_sheet(writer, "Periods", periods)
        _write_frame_sheet(writer, "Slices", slices)
        _write_frame_sheet(writer, "Trades", trades)
    workbook = load_workbook(path, read_only=True)
    expected = {
        "Generations": len(generations.columns),
        "Periods": len(periods.columns),
        "Slices": len(slices.columns),
        "Trades": len(trades.columns),
    }
    for name, columns in expected.items():
        if workbook[name].max_column != columns or columns <= 1:
            raise RuntimeError(f"invalid Excel table layout in {name}")


def plot_evolution(generations: pd.DataFrame, path: Path) -> None:
    generation = generations["generation"]
    figure, axes = plt.subplots(2, 2, figsize=(14, 9), constrained_layout=True)
    axes[0, 0].plot(generation, generations["best_fitness"], label="Best", color="#147d64")
    axes[0, 0].plot(generation, generations["p90_fitness"], label="P90", color="#2878b5")
    axes[0, 0].plot(generation, generations["median_fitness"], label="Median", color="#6b7280")
    axes[0, 0].axhline(0.0, color="black", linewidth=0.8)
    axes[0, 0].set_title("Generation fitness")
    axes[0, 0].legend()

    for column, label, color in (
        ("champion_mean_bps", "Mean", "#2878b5"),
        ("champion_median_bps", "Median", "#d97706"),
        ("champion_lcb_bps", "LCB 95%", "#b42318"),
    ):
        axes[0, 1].plot(generation, generations[column], label=label, color=color)
    axes[0, 1].axhline(0.0, color="black", linewidth=0.8)
    axes[0, 1].set_title("Champion costed expectancy")
    axes[0, 1].set_ylabel("bps per closed trade")
    axes[0, 1].legend()

    axes[1, 0].plot(generation, generations["champion_max_drawdown"] * 100.0, label="Drawdown %")
    axes[1, 0].plot(generation, generations["max_direction_share"] * 100.0, label="Max direction share %")
    axes[1, 0].plot(generation, generations["max_state_share"] * 100.0, label="Max state share %")
    axes[1, 0].plot(generation, generations["max_symbol_share"] * 100.0, label="Max symbol share %")
    axes[1, 0].plot(
        generation,
        generations["champion_constraint_violation"] * 100.0,
        label="Constraint violation x100",
        linestyle="--",
    )
    axes[1, 0].set_title("Champion risk and concentration")
    axes[1, 0].legend()

    axes[1, 1].plot(generation, generations["champion_trades"], label="Closed trades")
    axes[1, 1].plot(generation, generations["champion_turnover_fills"], label="Turnover fills")
    axes[1, 1].plot(generation, generations["flat_actions"], label="FLAT decisions", alpha=0.7)
    axes[1, 1].set_title("Champion activity")
    axes[1, 1].legend()
    for axis in axes.flat:
        axis.grid(alpha=0.25)
        axis.set_xlabel("Generation")
    figure.savefig(path, dpi=160)
    plt.close(figure)


def plot_regime_progress(generations: pd.DataFrame, path: Path) -> None:
    figure, axes = plt.subplots(1, 2, figsize=(15, 5), constrained_layout=True)
    for state in MARKET_STATES:
        mean_column = f"state_{state}_mean_bps"
        lcb_column = f"state_{state}_lcb_bps"
        if mean_column in generations:
            axes[0].plot(generations["generation"], generations[mean_column], label=state)
        if lcb_column in generations:
            axes[1].plot(generations["generation"], generations[lcb_column], label=state)
    for axis, title in zip(axes, ("Mean by market state", "LCB 95% by market state")):
        axis.axhline(0.0, color="black", linewidth=0.8)
        axis.set_title(title)
        axis.set_xlabel("Generation")
        axis.set_ylabel("bps per closed trade")
        axis.grid(alpha=0.25)
        axis.legend(fontsize=8)
    figure.savefig(path, dpi=160)
    plt.close(figure)


def plot_champion(champion: Champion, path: Path) -> None:
    figure, axes = plt.subplots(2, 1, figsize=(14, 8), constrained_layout=True)
    portfolio = champion.simulation.portfolio
    if not portfolio.empty:
        dates = pd.to_datetime(portfolio["timestamp"], unit="ms", utc=True)
        axes[0].plot(dates, portfolio["equity"], color="#147d64")
        axes[1].fill_between(dates, portfolio["drawdown"] * 100.0, 0.0, color="#b42318", alpha=0.35)
    axes[0].set_title("Development equity of selected champion")
    axes[1].set_title("Development drawdown")
    axes[1].set_ylabel("percent")
    for axis in axes:
        axis.grid(alpha=0.25)
    figure.savefig(path, dpi=160)
    plt.close(figure)


def save_generation_champion(output_dir: Path, champion: Champion) -> None:
    folder = output_dir / "generation_champions"
    folder.mkdir(parents=True, exist_ok=True)
    np.savez_compressed(
        folder / f"generation_{champion.generation:03d}.npz",
        genome=champion.genome,
        generation=np.asarray([champion.generation], dtype=np.int32),
        fitness=np.asarray([champion.score.fitness], dtype=np.float64),
    )
    (folder / f"generation_{champion.generation:03d}.json").write_text(
        json.dumps(dict(champion.score.metrics), indent=2, ensure_ascii=False, default=_json_default) + "\n",
        encoding="utf-8",
    )


def write_progress_snapshot(output_dir: Path, generations: pd.DataFrame) -> None:
    """Refresh lightweight human-readable progress after every generation."""
    generations.to_parquet(output_dir / "generation_progress.parquet", index=False)
    plot_evolution(generations, output_dir / "evolution_progress.png")
    plot_regime_progress(generations, output_dir / "market_state_progress_live.png")


def _json_default(value: Any) -> Any:
    if isinstance(value, (np.integer, np.floating)):
        return value.item()
    if isinstance(value, np.ndarray):
        return value.tolist()
    raise TypeError(type(value).__name__)


def write_evidence_bundle(
    output_dir: Path,
    *,
    result: EvolutionResult,
    manifest: Mapping[str, Any],
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    champion = result.champion
    np.savez_compressed(output_dir / "champion.npz", genome=champion.genome)
    result.generations.to_parquet(output_dir / "generation_metrics.parquet", index=False)
    champion.score.periods.to_parquet(output_dir / "period_metrics.parquet", index=False)
    champion.score.slices.to_parquet(output_dir / "slice_metrics.parquet", index=False)
    champion.simulation.ledger.to_parquet(output_dir / "trades.parquet", index=False)
    write_workbook(
        output_dir / "experiment_tables.xlsx",
        generations=result.generations,
        periods=champion.score.periods,
        slices=champion.score.slices,
        trades=champion.simulation.ledger,
    )
    plot_evolution(result.generations, output_dir / "evolution_overview.png")
    plot_regime_progress(result.generations, output_dir / "market_state_progress.png")
    plot_champion(champion, output_dir / "champion_equity_drawdown.png")
    (output_dir / "manifest.json").write_text(
        json.dumps(dict(manifest), indent=2, ensure_ascii=False, default=_json_default) + "\n",
        encoding="utf-8",
    )
    metrics = dict(champion.score.metrics)
    metrics.update(
        {
            "selected_generation": champion.generation,
            "orders_enabled": False,
            "promotion_authority": False,
        }
    )
    (output_dir / "summary.json").write_text(
        json.dumps(metrics, indent=2, ensure_ascii=False, default=_json_default) + "\n",
        encoding="utf-8",
    )
    report = [
        "# Genetic_Primus experiment",
        "",
        f"Status: `{manifest['status']}`.",
        "",
        "Research only. This artifact cannot authorize paper or live orders.",
        "",
        "## Selected development champion",
        "",
        f"- Generation: {champion.generation}",
        f"- Fitness: {champion.score.fitness:.4f}",
        f"- Closed trades: {metrics['trades']}",
        f"- Mean net: {metrics['mean_bps']} bps",
        f"- Median net: {metrics['median_bps']} bps",
        f"- LCB 95% net: {metrics['lcb_bps']} bps",
        f"- Max drawdown: {metrics['max_drawdown']:.4f}",
        f"- Failed development gates: {', '.join(metrics['failed_gates']) or 'none'}",
        "",
        "## Evidence",
        "",
        "- `evolution_overview.png`: fitness, expectancy, drawdown and activity by generation.",
        "- `market_state_progress.png`: mean and LCB by causal market state.",
        "- `champion_equity_drawdown.png`: champion equity and drawdown.",
        "- `experiment_tables.xlsx`: separate Generation, Period, Slice and Trade worksheets.",
        "- `generation_champions/`: immutable best genome and metrics for every generation.",
        "",
        "Validation/OOS/sanity were not opened by this development runner.",
        "",
    ]
    (output_dir / "report.md").write_text("\n".join(report), encoding="utf-8")
