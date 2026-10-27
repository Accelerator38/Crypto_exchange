from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Mapping

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import TwoSlopeNorm
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .fitness import FitnessResult


def strict_development_status(failed_gates: tuple[str, ...]) -> str:
    if not failed_gates:
        return "DEVELOPMENT_PASSED_STRICT_GATES"
    if any(
        failure.startswith(
            (
                "nonpositive_market_state_lcb:",
                "nonpositive_direction_lcb:",
                "nonpositive_symbol_lcb:",
            )
        )
        for failure in failed_gates
    ):
        return "DEVELOPMENT_REJECTED_SLICE_GATES"
    return "DEVELOPMENT_REJECTED_STRICT_GATES"


def build_freeze_manifest(
    *,
    experiment_manifest: Mapping[str, Any],
    experiment_manifest_sha256: str,
    champion_file_sha256: str,
    champion_genome_sha256: str,
    scaler_sha256: str,
    score: FitnessResult,
) -> dict[str, Any]:
    failures = list(score.failed_gates)
    status = strict_development_status(score.failed_gates)
    validation_allowed = status == "DEVELOPMENT_PASSED_STRICT_GATES"
    return {
        "schema_version": "exia.genetic_primus_freeze.v1",
        "agent_id": "Genetic_Primus",
        "generated_at": datetime.now(UTC).isoformat(),
        "status": status,
        "experiment_manifest_sha256": experiment_manifest_sha256,
        "config_sha256": experiment_manifest["config_sha256"],
        "code_sha256": experiment_manifest["code_sha256"],
        "dataset_sha256": experiment_manifest["dataset_sha256"],
        "champion_file_sha256": champion_file_sha256,
        "champion_genome_sha256": champion_genome_sha256,
        "scaler_sha256": scaler_sha256,
        "split": "development",
        "cost_profile": "stress_round_trip",
        "metrics": dict(score.metrics),
        "failed_gates": failures,
        "validation_allowed": validation_allowed,
        "validation_opened": False,
        "oos_opened": False,
        "sanity_opened": False,
        "paper_allowed": False,
        "live_allowed": False,
        "orders_enabled": False,
        "promotion_authority": False,
    }


def _write_sheet(writer: pd.ExcelWriter, name: str, frame: pd.DataFrame) -> None:
    frame.to_excel(writer, sheet_name=name, index=False)
    sheet = writer.book[name]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        values = [str(cell.value) if cell.value is not None else "" for cell in column[:200]]
        sheet.column_dimensions[column[0].column_letter].width = min(
            48, max(10, max(len(value) for value in values) + 2)
        )


def write_strict_workbook(
    path: Path,
    *,
    manifest: Mapping[str, Any],
    periods: pd.DataFrame,
    slices: pd.DataFrame,
) -> None:
    metrics = manifest["metrics"]
    summary = pd.DataFrame(
        [
            {
                "status": manifest["status"],
                "validation_allowed": manifest["validation_allowed"],
                "trades": metrics["trades"],
                "mean_bps": metrics["mean_bps"],
                "median_bps": metrics["median_bps"],
                "lcb_bps": metrics["lcb_bps"],
                "max_drawdown": metrics["max_drawdown"],
                "fitness": metrics["fitness"],
                "champion_genome_sha256": manifest["champion_genome_sha256"],
            }
        ]
    )
    failures = pd.DataFrame(
        {"failed_gate": list(manifest["failed_gates"]) or ["none"]}
    )
    frames = {
        "Summary": summary,
        "Periods": periods,
        "Slices": slices,
        "Failed Gates": failures,
    }
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in frames.items():
            _write_sheet(writer, name, frame)

    workbook = load_workbook(path, read_only=False)
    for name, frame in frames.items():
        sheet = workbook[name]
        if sheet.max_column != len(frame.columns) or sheet.max_column <= 0:
            raise RuntimeError(f"invalid Excel table layout in {name}")
        if sheet.merged_cells.ranges:
            raise RuntimeError(f"merged cells are forbidden in {name}")
    workbook.close()


def plot_strict_slices(
    slices: pd.DataFrame,
    path: Path,
    *,
    title: str = "Genetic_Primus strict development slices",
) -> None:
    metrics = ("mean_bps", "median_bps", "lcb_bps")
    labels = ("Mean", "Median", "LCB 95%")
    values = slices.loc[:, metrics].astype("float64").to_numpy()
    finite = np.abs(values[np.isfinite(values)])
    limit = max(1.0, float(np.quantile(finite, 0.95))) if len(finite) else 1.0
    norm = TwoSlopeNorm(vmin=-limit, vcenter=0.0, vmax=limit)
    figure, axes = plt.subplots(1, 3, figsize=(18, 8), constrained_layout=True)
    for axis, scope in zip(axes, ("market_state", "direction", "symbol")):
        subset = slices.loc[slices["scope"] == scope].reset_index(drop=True)
        matrix = subset.loc[:, metrics].astype("float64").to_numpy()
        image = axis.imshow(matrix, cmap="RdYlGn", norm=norm, aspect="auto")
        axis.set_xticks(range(len(labels)), labels=labels)
        axis.set_yticks(range(len(subset)), labels=subset["value"].astype(str))
        axis.set_title(scope.replace("_", " ").title())
        for row in range(len(subset)):
            for column in range(len(metrics)):
                value = matrix[row, column]
                text = "n/a" if not np.isfinite(value) else f"{value:.1f}"
                axis.text(column, row, text, ha="center", va="center", fontsize=8)
        axis.set_xlabel("net bps per closed trade")
    figure.colorbar(image, ax=axes, shrink=0.75, label="costed bps")
    figure.suptitle(title)
    figure.savefig(path, dpi=170)
    plt.close(figure)


def write_strict_report(
    path: Path,
    *,
    manifest: Mapping[str, Any],
    slices: pd.DataFrame,
) -> None:
    metrics = manifest["metrics"]
    failed = "\n".join(f"- `{value}`" for value in manifest["failed_gates"])
    if not failed:
        failed = "- none"
    weakest = slices.sort_values("lcb_bps", na_position="first").head(8)
    weak_rows = "\n".join(
        f"- `{row.scope}:{row.value}`: trades={int(row.trades)}, "
        f"mean={row.mean_bps:.3f}, median={row.median_bps:.3f}, LCB={row.lcb_bps:.3f} bps"
        for row in weakest.itertuples(index=False)
    )
    report = f"""# Genetic_Primus V1 strict development freeze

Status: `{manifest['status']}`.

Validation allowed: `{str(manifest['validation_allowed']).lower()}`. Validation, OOS,
sanity, paper and live were not opened. Evaluation uses the stress round-trip cost.

## Overall development result

- Closed trades: {int(metrics['trades'])}
- Mean: {metrics['mean_bps']:.6f} bps
- Median: {metrics['median_bps']:.6f} bps
- LCB 95%: {metrics['lcb_bps']:.6f} bps
- Max drawdown: {metrics['max_drawdown']:.6%}
- Supported months: {int(metrics['supported_periods'])}/{int(metrics['total_periods'])}

## Failed strict gates

{failed}

## Weakest supported slices

{weak_rows}

## Immutable identity

- Experiment manifest SHA: `{manifest['experiment_manifest_sha256']}`
- Champion file SHA: `{manifest['champion_file_sha256']}`
- Raw genome SHA: `{manifest['champion_genome_sha256']}`
- Scaler SHA: `{manifest['scaler_sha256']}`
- Dataset SHA: `{manifest['dataset_sha256']}`

This freeze artifact has no promotion authority.
"""
    path.write_text(report, encoding="utf-8")


def write_json(path: Path, payload: Mapping[str, Any]) -> None:
    path.write_text(
        json.dumps(dict(payload), indent=2, ensure_ascii=False, default=_json_default) + "\n",
        encoding="utf-8",
    )


def _json_default(value: Any) -> Any:
    if isinstance(value, (np.integer, np.floating)):
        return value.item()
    if isinstance(value, np.ndarray):
        return value.tolist()
    raise TypeError(type(value).__name__)
