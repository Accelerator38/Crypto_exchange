from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
RUNNER_PATH = ROOT / "tools/compare_genetic_primus_models_v1.py"


def _load_runner():
    spec = importlib.util.spec_from_file_location("primus_comparison_test", RUNNER_PATH)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_common_summary_keeps_fields_in_separate_columns() -> None:
    runner = _load_runner()
    ledger = pd.DataFrame(
        {
            "net_bps": [10.0, 20.0, -5.0],
            "signal_timestamp": [0, 14_400_000, 28_800_000],
            "fills": [2, 2, 2],
            "direction": ["LONG", "LONG", "SHORT"],
            "symbol": ["BTC/USDT", "ETH/USDT", "BTC/USDT"],
        }
    )
    result = runner.summarize_ledger(
        ledger,
        model_id="TEST",
        split="development",
        max_drawdown=0.1,
    )
    assert result["trades"] == 3
    assert result["fills"] == 6
    assert result["mean_bps"] == 25.0 / 3.0
    assert result["median_bps"] == 10.0
    assert result["max_direction_share"] == 2.0 / 3.0


def test_workbook_has_no_merged_or_one_cell_tables(tmp_path: Path) -> None:
    runner = _load_runner()
    leaderboard = pd.DataFrame(
        [
            {
                "model_id": "Genetic_Primus_V3",
                "model_family": "neurogenetic",
                "split": "development",
                "mean_bps": 1.0,
                "median_bps": 0.5,
                "lcb_bps": 0.1,
            },
            {
                "model_id": "MQ_EMA12_48_BOTH",
                "model_family": "standard_market_quorum",
                "split": "development",
                "mean_bps": -1.0,
                "median_bps": -0.5,
                "lcb_bps": -2.0,
            },
            {
                "model_id": "player_one",
                "model_family": "legacy_ensemble_reference",
                "split": "development",
                "mean_bps": -1.0,
                "median_bps": -0.5,
                "lcb_bps": -2.0,
            },
        ]
    )
    methodology = pd.DataFrame([{"field": "cost_bps", "value": 16.0}])
    path = tmp_path / "comparison.xlsx"
    runner.write_workbook(path, leaderboard=leaderboard, methodology=methodology)

    workbook = load_workbook(path, read_only=False)
    for sheet in workbook.worksheets:
        assert sheet.max_column > 1
        assert not sheet.merged_cells.ranges
    assert workbook["Common Leaderboard"]["A1"].value == "model_id"
    assert workbook["Common Leaderboard"]["B1"].value == "model_family"
    workbook.close()
