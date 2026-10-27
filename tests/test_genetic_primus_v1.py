from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from exia.genetic_primus.contracts import FEATURE_NAMES, load_primus_config
from exia.genetic_primus.evolution import Champion, EvolutionResult, run_evolution
from exia.genetic_primus.features import RobustScaler, build_feature_frame
from exia.genetic_primus.fitness import score_simulation
from exia.genetic_primus.freeze import build_freeze_manifest, strict_development_status
from exia.genetic_primus.labels import build_forward_labels
from exia.genetic_primus.policy import GeneticPrimusPolicy, NetworkShape
from exia.genetic_primus.reporting import write_evidence_bundle
from exia.genetic_primus.simulator import PrimusSimulation, simulate_policy


ROOT = Path(__file__).resolve().parents[1]
SYMBOLS = (
    "ADA/USDT",
    "BNB/USDT",
    "BTC/USDT",
    "DOGE/USDT",
    "ETH/USDT",
    "LINK/USDT",
    "SOL/USDT",
    "XRP/USDT",
)


def synthetic_panel(*, bars: int = 180) -> pd.DataFrame:
    timestamps = 1_640_995_200_000 + np.arange(bars, dtype=np.int64) * 14_400_000
    rows: list[dict[str, object]] = []
    for symbol_index, symbol in enumerate(SYMBOLS):
        base = 100.0 + symbol_index * 10.0
        for index, timestamp in enumerate(timestamps):
            trend = 1.0 + 0.0008 * index
            wave = 1.0 + 0.003 * np.sin(index / 8.0 + symbol_index)
            close = base * trend * wave
            opened = close * (1.0 - 0.0005 * np.cos(index / 5.0))
            rows.append(
                {
                    "timestamp": int(timestamp),
                    "open": opened,
                    "high": max(opened, close) * 1.002,
                    "low": min(opened, close) * 0.998,
                    "close": close,
                    "volume": 1_000.0 + index * 2.0 + symbol_index * 25.0,
                    "symbol": symbol,
                }
            )
    return pd.DataFrame(rows).sort_values(["timestamp", "symbol"]).reset_index(drop=True)


def constant_policy(target: int) -> GeneticPrimusPolicy:
    shape = NetworkShape()
    genome = np.zeros(shape.genome_size, dtype=np.float32)
    weights = list(GeneticPrimusPolicy.unpack(genome, shape))
    weights[-1][{-1: 0, 0: 1, 1: 2}[target]] = 1.0
    return GeneticPrimusPolicy(GeneticPrimusPolicy.pack(tuple(weights)))


def test_config_is_fixed_4h_research_only() -> None:
    config = load_primus_config(ROOT / "configs/genetic_primus_v1.json", root=ROOT)
    assert config.payload["agent_id"] == "Genetic_Primus"
    assert config.dataset["timeframe_minutes"] == 240
    assert config.network["inputs"] == len(FEATURE_NAMES) == 30
    assert config.payload["safety"] == {
        "paper_allowed": False,
        "live_allowed": False,
        "orders_enabled": False,
        "promotion_authority": False,
    }


def test_features_are_causal_and_have_five_market_states() -> None:
    panel = synthetic_panel()
    original = build_feature_frame(panel)
    changed = panel.copy()
    final_timestamp = int(changed["timestamp"].max())
    changed.loc[changed["timestamp"] == final_timestamp, "close"] *= 5.0
    changed.loc[changed["timestamp"] == final_timestamp, "high"] *= 5.0
    rebuilt = build_feature_frame(changed)
    previous = original["timestamp"] < final_timestamp
    pd.testing.assert_frame_equal(
        original.loc[previous, list(FEATURE_NAMES[:25])].reset_index(drop=True),
        rebuilt.loc[previous, list(FEATURE_NAMES[:25])].reset_index(drop=True),
    )
    assert original["feature_supported"].sum() > 0
    assert set(original["market_state"].dropna().unique()) <= {
        "trend_up", "trend_down", "quiet_range", "volatile_range", "transition"
    }


def test_forward_labels_use_next_open_and_nonzero_costs() -> None:
    panel = synthetic_panel(bars=80)
    labels = build_forward_labels(
        panel,
        round_trip_cost_bps=16.0,
        minimum_edge_bps=0.0,
    )
    btc = labels.loc[labels["symbol"] == "BTC/USDT"].reset_index(drop=True)
    expected_gross = (btc.loc[2, "open"] / btc.loc[1, "open"] - 1.0) * 10_000.0
    assert btc.loc[0, "long_net_h1_bps"] == expected_gross - 16.0
    assert btc.loc[0, "short_net_h1_bps"] == -expected_gross - 16.0
    assert not btc.iloc[-7:]["label_supported"].all()


def test_target_position_simulator_executes_at_next_open() -> None:
    features = build_feature_frame(synthetic_panel())
    supported = features.loc[features["feature_supported"]]
    scaler = RobustScaler.fit(supported)
    start = int(features["timestamp"].min())
    end = int(features["timestamp"].max()) + 14_400_000
    result = simulate_policy(
        features,
        policy=constant_policy(1),
        scaler=scaler,
        start_timestamp=start,
        end_timestamp=end,
        round_trip_cost_bps=16.0,
    )
    first_supported = (
        features.loc[features["feature_supported"]]
        .groupby("symbol", sort=True)["timestamp"]
        .min()
    )
    for row in result.ledger.itertuples(index=False):
        assert row.entry_timestamp == first_supported[row.symbol] + 14_400_000
        assert row.cost_bps == 16.0
        assert row.fills == 2
    assert len(result.ledger) == len(SYMBOLS)
    assert result.turnover_fills == len(SYMBOLS) * 2


def test_zero_trade_month_is_not_dropped_from_fitness() -> None:
    january = int(pd.Timestamp("2022-01-10T00:00:00Z").timestamp() * 1000)
    march = int(pd.Timestamp("2022-03-10T00:00:00Z").timestamp() * 1000)
    ledger = pd.DataFrame(
        [
            {
                "symbol": "BTC/USDT",
                "direction": "LONG",
                "market_state": "trend_up",
                "signal_timestamp": january,
                "entry_timestamp": january,
                "exit_timestamp": january + 14_400_000,
                "entry_price": 100.0,
                "exit_price": 101.0,
                "holding_bars": 1,
                "gross_bps": 100.0,
                "cost_bps": 16.0,
                "net_bps": 84.0,
                "fills": 2,
                "exit_reason": "target_change",
            }
        ]
    )
    portfolio = pd.DataFrame(
        {
            "timestamp": [january, march],
            "net_return": [0.0, 0.0],
            "equity": [1.0, 1.0],
            "drawdown": [0.0, 0.0],
        }
    )
    simulation = PrimusSimulation(
        ledger=ledger,
        portfolio=portfolio,
        action_counts={"SHORT": 0, "FLAT": 1, "LONG": 1},
        turnover_fills=2,
    )
    config = {
        "weights": {"mean": 0.2, "median": 0.3, "lcb": 0.5},
        "bootstrap_samples": 100,
        "bootstrap_alpha": 0.05,
        "minimum_period_trades": 1,
        "minimum_split_trades": 1,
        "minimum_slice_trades": 1,
        "minimum_supported_period_rate": 0.0,
        "max_drawdown": 1.0,
        "max_direction_share": 1.0,
        "max_state_share": 1.0,
        "max_symbol_share": 1.0,
    }
    score = score_simulation(simulation, config=config, key="zero-month")
    assert score.periods["period"].tolist() == ["2022-01", "2022-02", "2022-03"]
    february = score.periods.loc[score.periods["period"] == "2022-02"].iloc[0]
    assert february["trades"] == 0
    assert bool(february["supported"]) is False


def test_evolution_resumes_from_atomic_checkpoint(tmp_path: Path) -> None:
    features = build_feature_frame(synthetic_panel(bars=100))
    scaler = RobustScaler.fit(features.loc[features["feature_supported"]])
    fitness_config = {
        "weights": {"mean": 0.2, "median": 0.3, "lcb": 0.5},
        "bootstrap_samples": 50,
        "bootstrap_alpha": 0.05,
        "minimum_period_trades": 1,
        "minimum_split_trades": 1,
        "minimum_slice_trades": 1,
        "minimum_supported_period_rate": 0.0,
        "max_drawdown": 1.0,
        "max_direction_share": 1.0,
        "max_state_share": 1.0,
        "max_symbol_share": 1.0,
    }
    evolution_config = {
        "seed": 7,
        "workers": 1,
        "population_size": 4,
        "generations": 2,
        "elite_count": 1,
        "tournament_size": 2,
        "crossover_rate": 0.7,
        "mutation_rate": 0.15,
        "mutation_sigma": 0.08,
    }
    checkpoint = tmp_path / "checkpoint.npz"

    def interrupt(_champion, _history) -> None:
        raise RuntimeError("simulated power interruption")

    with pytest.raises(RuntimeError, match="power interruption"):
        run_evolution(
            features,
            scaler=scaler,
            start_timestamp=int(features["timestamp"].min()),
            end_timestamp=int(features["timestamp"].max()) + 14_400_000,
            round_trip_cost_bps=16.0,
            fitness_config=fitness_config,
            evolution_config=evolution_config,
            checkpoint_path=checkpoint,
            run_fingerprint="test-fingerprint",
            on_generation=interrupt,
        )
    assert checkpoint.is_file()
    resumed = run_evolution(
        features,
        scaler=scaler,
        start_timestamp=int(features["timestamp"].min()),
        end_timestamp=int(features["timestamp"].max()) + 14_400_000,
        round_trip_cost_bps=16.0,
        fitness_config=fitness_config,
        evolution_config=evolution_config,
        checkpoint_path=checkpoint,
        run_fingerprint="test-fingerprint",
        resume=True,
    )
    assert resumed.generations["generation"].tolist() == [0, 1]
    with pytest.raises(ValueError, match="fingerprint mismatch"):
        run_evolution(
            features,
            scaler=scaler,
            start_timestamp=int(features["timestamp"].min()),
            end_timestamp=int(features["timestamp"].max()) + 14_400_000,
            round_trip_cost_bps=16.0,
            fitness_config=fitness_config,
            evolution_config=evolution_config,
            checkpoint_path=checkpoint,
            run_fingerprint="wrong",
            resume=True,
        )


def test_evidence_bundle_has_charts_and_separate_excel_cells(tmp_path: Path) -> None:
    features = build_feature_frame(synthetic_panel())
    scaler = RobustScaler.fit(features.loc[features["feature_supported"]])
    simulation = simulate_policy(
        features,
        policy=constant_policy(1),
        scaler=scaler,
        start_timestamp=int(features["timestamp"].min()),
        end_timestamp=int(features["timestamp"].max()) + 14_400_000,
        round_trip_cost_bps=16.0,
    )
    fitness_config = {
        "weights": {"mean": 0.2, "median": 0.3, "lcb": 0.5},
        "bootstrap_samples": 100,
        "bootstrap_alpha": 0.05,
        "minimum_period_trades": 1,
        "minimum_split_trades": 1,
        "minimum_slice_trades": 1,
        "minimum_supported_period_rate": 0.0,
        "max_drawdown": 1.0,
        "max_direction_share": 1.0,
        "max_state_share": 1.0,
        "max_symbol_share": 1.0,
    }
    score = score_simulation(simulation, config=fitness_config, key="test")
    champion = Champion(0, constant_policy(1).genome, simulation, score)
    generation = {
        "generation": 0,
        "best_fitness": score.fitness,
        "best_selection_score": score.selection_score,
        "median_fitness": score.fitness,
        "p90_fitness": score.fitness,
        "champion_trades": len(simulation.ledger),
        "champion_mean_bps": score.metrics["mean_bps"],
        "champion_median_bps": score.metrics["median_bps"],
        "champion_lcb_bps": score.metrics["lcb_bps"],
        "champion_max_drawdown": score.metrics["max_drawdown"],
        "champion_constraint_violation": score.constraint_violation,
        "champion_turnover_fills": score.metrics["turnover_fills"],
        "supported_period_rate": score.metrics["supported_period_rate"],
        "max_direction_share": score.metrics["max_direction_share"],
        "max_state_share": score.metrics["max_state_share"],
        "max_symbol_share": score.metrics["max_symbol_share"],
        "short_actions": 0,
        "flat_actions": 0,
        "long_actions": 1,
        "failed_gates": "",
    }
    for state in ("trend_up", "trend_down", "quiet_range", "volatile_range", "transition"):
        generation[f"state_{state}_trades"] = 0
        generation[f"state_{state}_mean_bps"] = np.nan
        generation[f"state_{state}_median_bps"] = np.nan
        generation[f"state_{state}_lcb_bps"] = np.nan
    evolution = EvolutionResult(champion, pd.DataFrame([generation]), (champion,))
    write_evidence_bundle(
        tmp_path,
        result=evolution,
        manifest={"status": "ENGINEERING_SMOKE", "orders_enabled": False, "promotion_authority": False},
    )
    for name in (
        "evolution_overview.png",
        "market_state_progress.png",
        "champion_equity_drawdown.png",
        "experiment_tables.xlsx",
        "manifest.json",
        "report.md",
    ):
        assert (tmp_path / name).is_file()
    workbook = load_workbook(tmp_path / "experiment_tables.xlsx")
    assert workbook["Generations"].max_column == len(evolution.generations.columns)
    assert workbook["Generations"].max_column > 1
    assert not workbook["Generations"].merged_cells.ranges
    manifest = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["orders_enabled"] is False


def test_strict_freeze_rejects_slice_gate_and_keeps_all_authority_disabled() -> None:
    features = build_feature_frame(synthetic_panel())
    scaler = RobustScaler.fit(features.loc[features["feature_supported"]])
    simulation = simulate_policy(
        features,
        policy=constant_policy(1),
        scaler=scaler,
        start_timestamp=int(features["timestamp"].min()),
        end_timestamp=int(features["timestamp"].max()) + 14_400_000,
        round_trip_cost_bps=16.0,
    )
    fitness_config = {
        "weights": {"mean": 0.2, "median": 0.3, "lcb": 0.5},
        "bootstrap_samples": 50,
        "bootstrap_alpha": 0.05,
        "minimum_period_trades": 1,
        "minimum_split_trades": 1,
        "minimum_slice_trades": 1,
        "minimum_supported_period_rate": 0.0,
        "max_drawdown": 1.0,
        "max_direction_share": 1.0,
        "max_state_share": 1.0,
        "max_symbol_share": 1.0,
    }
    score = score_simulation(
        simulation,
        config=fitness_config,
        key="strict-freeze-test",
        final_gate=True,
    )
    assert any(gate.startswith("nonpositive_") for gate in score.failed_gates)
    freeze = build_freeze_manifest(
        experiment_manifest={
            "config_sha256": "config",
            "code_sha256": "code",
            "dataset_sha256": "dataset",
        },
        experiment_manifest_sha256="manifest",
        champion_file_sha256="champion",
        champion_genome_sha256="genome",
        scaler_sha256="scaler",
        score=score,
    )
    assert strict_development_status(score.failed_gates).startswith("DEVELOPMENT_REJECTED")
    assert freeze["validation_allowed"] is False
    assert freeze["validation_opened"] is False
    assert freeze["paper_allowed"] is False
    assert freeze["live_allowed"] is False
    assert freeze["orders_enabled"] is False
    assert freeze["promotion_authority"] is False
